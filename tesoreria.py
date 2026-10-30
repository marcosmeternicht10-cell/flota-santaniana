"""
tesoreria.py — Tesorería: flujo de caja, pagos y cobros programados, cheques
y conciliación bancaria (La Santaniana)

Lo que tienen los programas de tesorería serios (Agicap, Kyriba, Sage,
Orama…), llevado a lo que necesita la empresa, encima de las Cajas:

  FLUJO DE CAJA   El saldo de hoy de todas las cajas y cuentas, más lo que
                  viene (pagos y cobros programados, cheques por cobrar y por
                  debitar, y si se quiere el gasto habitual de cada semana),
                  semana por semana durante 13 semanas. Avisa en qué semana
                  la plata no alcanza o baja del mínimo.
  PAGOS Y COBROS  Lo que hay que pagar o cobrar y cuándo: vencimientos de
                  proveedores, sueldos, alquiler, cuotas… Se repiten solos
                  (semanal, quincenal, mensual). Los pagos grandes necesitan
                  la aprobación de administración antes de pagarse. Al
                  pagarlo se carga el movimiento en la caja.
  CHEQUES         La cartera de cheques recibidos y la chequera de emitidos,
                  comunes y de pago diferido (Ley 805/96: la fecha de pago no
                  puede pasar de 180 días desde la emisión y hay 30 días
                  después de esa fecha para presentarlo). Avisa los que
                  vencen, y al cobrarse o debitarse se carga el movimiento.
  CONCILIACIÓN    Se sube el extracto del banco (Excel o CSV) y se cruza
                  solo con lo cargado en esa cuenta: mismo monto, fecha
                  cercana. Lo que queda suelto se engancha a mano o se carga
                  (comisiones, intereses) con un clic.

  init_tesoreria()   tablas y columnas nuevas
  bp_teso            /api/teso/...
"""

import csv
import io
import json
import re
from datetime import date, datetime, timedelta

from flask import Blueprint, request, jsonify, session

from db_compat import get_connection, USE_POSTGRES
from hora_local import hoy as _hoy, ahora_iso

bp_teso = Blueprint("tesoreria", __name__)

PK = "SERIAL PRIMARY KEY" if USE_POSTGRES else "INTEGER PRIMARY KEY AUTOINCREMENT"

RECURRENCIAS = {"": "Una sola vez", "semanal": "Cada semana", "quincenal": "Cada 15 días", "mensual": "Cada mes",
                "bimestral": "Cada 2 meses", "anual": "Cada año"}
CHEQUE_DIFERIDO_MAX_DIAS = 180     # Ley 805/96, art. 1.696 CC: la fecha de pago no puede pasar de 180 días
CHEQUE_PRESENTACION_DIAS = 30      # Ley 805/96, art. 1.726 CC: 30 días desde la fecha de pago para presentarlo
ESTADOS_CHEQUE = {
    "recibido": {"cartera": "En cartera", "depositado": "Depositado", "cobrado": "Cobrado",
                 "endosado": "Endosado a un tercero", "rechazado": "Rechazado"},
    "emitido": {"emitido": "Emitido (sin entregar)", "entregado": "Entregado", "debitado": "Debitado",
                "anulado": "Anulado", "rechazado": "Rechazado"},
}
BANCOS = ["Itaú", "Continental", "Banco Basa", "Visión Banco", "Banco Familiar", "Sudameris", "Banco GNB",
          "Banco Atlas", "Banco Nacional de Fomento", "Interfisa", "Ueno Bank", "Bancop", "Banco Río",
          "Zeta Banco", "Citibank", "Banco do Brasil", "Banco de la Nación Argentina", "Solar Banco", "FIC"]


def _rollback(conn):
    try:
        conn.rollback()
    except Exception:
        pass


def _quien():
    return session.get("nombre") or session.get("usuario") or "?"


def _txt(v, n):
    return re.sub(r"\s+", " ", str(v or "")).strip()[:n]


def _fecha(v):
    s = str(v or "").strip()[:10]
    try:
        return date.fromisoformat(s).isoformat()
    except ValueError:
        pass
    m = re.match(r"(\d{1,2})[/.-](\d{1,2})[/.-](\d{2,4})", s)
    if m:
        d, mth, y = int(m.group(1)), int(m.group(2)), int(m.group(3))
        y = y + 2000 if y < 100 else y
        try:
            return date(y, mth, d).isoformat()
        except ValueError:
            return ""
    return ""


def _num(v):
    import cajas
    return cajas._num(v)


def _columna(conn, tabla, col, tipo):
    try:
        if USE_POSTGRES:
            conn.execute(f"ALTER TABLE {tabla} ADD COLUMN IF NOT EXISTS {col} {tipo}")
        else:
            ya = {r["name"] for r in conn.execute(f"PRAGMA table_info({tabla})").fetchall()}
            if col not in ya:
                conn.execute(f"ALTER TABLE {tabla} ADD COLUMN {col} {tipo}")
        conn.commit()
    except Exception:
        _rollback(conn)


def init_tesoreria():
    conn = get_connection()
    try:
        for col, tipo in (("tipo", "TEXT DEFAULT 'caja'"), ("banco", "TEXT DEFAULT ''"), ("numero_cuenta", "TEXT DEFAULT ''"),
                          ("saldo_minimo", "REAL DEFAULT 0")):
            _columna(conn, "cajas", col, tipo)
        _columna(conn, "caja_movs", "extracto_id", "INTEGER DEFAULT 0")
        for sql in (
            f"""CREATE TABLE IF NOT EXISTS teso_programados (
                id {PK},
                tipo TEXT NOT NULL,
                caja_id INTEGER DEFAULT 0,
                concepto_id INTEGER DEFAULT 0,
                concepto TEXT DEFAULT '',
                descripcion TEXT DEFAULT '',
                contraparte TEXT DEFAULT '',
                ruc TEXT DEFAULT '',
                doc_numero TEXT DEFAULT '',
                monto REAL NOT NULL,
                fecha TEXT NOT NULL,
                recurrencia TEXT DEFAULT '',
                hasta TEXT DEFAULT '',
                medio TEXT DEFAULT 'transferencia',
                prioridad TEXT DEFAULT 'normal',
                aprobacion TEXT DEFAULT 'no_requiere',
                aprobado_por TEXT DEFAULT '',
                aprobado_el TEXT DEFAULT '',
                estado TEXT DEFAULT 'activo',
                nota TEXT DEFAULT '',
                creado_por TEXT DEFAULT '',
                creado_el TEXT DEFAULT '')""",
            f"""CREATE TABLE IF NOT EXISTS teso_prog_hechos (
                id {PK},
                prog_id INTEGER NOT NULL,
                fecha TEXT NOT NULL,
                accion TEXT DEFAULT 'pagado',
                mov_id INTEGER DEFAULT 0,
                monto REAL DEFAULT 0,
                hecho_por TEXT DEFAULT '',
                hecho_el TEXT DEFAULT '',
                UNIQUE (prog_id, fecha))""",
            f"""CREATE TABLE IF NOT EXISTS teso_cheques (
                id {PK},
                sentido TEXT NOT NULL,
                tipo TEXT DEFAULT 'diferido',
                banco TEXT DEFAULT '',
                numero TEXT DEFAULT '',
                persona TEXT DEFAULT '',
                ruc TEXT DEFAULT '',
                monto REAL NOT NULL,
                fecha_emision TEXT NOT NULL,
                fecha_pago TEXT NOT NULL,
                estado TEXT NOT NULL,
                caja_id INTEGER DEFAULT 0,
                mov_id INTEGER DEFAULT 0,
                concepto TEXT DEFAULT '',
                nota TEXT DEFAULT '',
                historial TEXT DEFAULT '[]',
                creado_por TEXT DEFAULT '',
                creado_el TEXT DEFAULT '')""",
            f"""CREATE TABLE IF NOT EXISTS teso_extracto (
                id {PK},
                caja_id INTEGER NOT NULL,
                fecha TEXT NOT NULL,
                descripcion TEXT DEFAULT '',
                referencia TEXT DEFAULT '',
                monto REAL NOT NULL,
                saldo REAL,
                lote TEXT DEFAULT '',
                mov_id INTEGER DEFAULT 0,
                estado TEXT DEFAULT 'pendiente',
                como TEXT DEFAULT '',
                importado_por TEXT DEFAULT '',
                importado_el TEXT DEFAULT '')""",
            """CREATE TABLE IF NOT EXISTS teso_config (
                clave TEXT PRIMARY KEY,
                valor TEXT DEFAULT '')""",
            "CREATE INDEX IF NOT EXISTS ix_tprog_fecha ON teso_programados (fecha)",
            "CREATE INDEX IF NOT EXISTS ix_tcheq_fecha ON teso_cheques (fecha_pago)",
            "CREATE INDEX IF NOT EXISTS ix_text_caja ON teso_extracto (caja_id, fecha)",
        ):
            conn.execute(sql)
        conn.commit()
    except Exception as e:
        _rollback(conn)
        print(f"[tesoreria] no se pudieron crear las tablas: {e}")
    finally:
        conn.close()


def config():
    conn = get_connection()
    try:
        d = {r["clave"]: r["valor"] for r in conn.execute("SELECT clave, valor FROM teso_config").fetchall()}
    except Exception:
        _rollback(conn)
        d = {}
    finally:
        conn.close()
    return {"umbral_aprobacion": float(d.get("umbral_aprobacion") or 5000000)}


def guardar_config(d):
    conn = get_connection()
    try:
        if "umbral_aprobacion" in d:
            v = str(max(0, round(_num(d["umbral_aprobacion"]))))
            if conn.execute("SELECT 1 FROM teso_config WHERE clave='umbral_aprobacion'").fetchone():
                conn.execute("UPDATE teso_config SET valor=? WHERE clave='umbral_aprobacion'", (v,))
            else:
                conn.execute("INSERT INTO teso_config (clave, valor) VALUES ('umbral_aprobacion', ?)", (v,))
        conn.commit()
    finally:
        conn.close()
    return config()


# ════════════════════════════════════════════════════════════════════════════
# PAGOS Y COBROS PROGRAMADOS
# ════════════════════════════════════════════════════════════════════════════

def _sumar_periodo(f, rec, n=1):
    if rec == "semanal":
        return f + timedelta(days=7 * n)
    if rec == "quincenal":
        return f + timedelta(days=15 * n)
    meses = {"mensual": 1, "bimestral": 2, "anual": 12}.get(rec)
    if not meses:
        return None
    m = f.month - 1 + meses * n
    y, m = f.year + m // 12, m % 12 + 1
    import calendar
    return date(y, m, min(f.day, calendar.monthrange(y, m)[1]))


def ocurrencias(p, desde, hasta, hechos):
    """Las fechas en que toca un programado entre desde y hasta (sin las ya
    hechas). Las vencidas sin pagar se devuelven también (desde la primera)."""
    f0 = date.fromisoformat(p["fecha"])
    fin = date.fromisoformat(hasta)
    tope = date.fromisoformat(p["hasta"]) if p.get("hasta") else None
    out = []
    k = 0
    f = f0
    while f and f <= fin and k < 400:
        if tope and f > tope:
            break
        fs = f.isoformat()
        if fs not in hechos:
            out.append(fs)
        if not p.get("recurrencia"):
            break
        k += 1
        f = _sumar_periodo(f0, p["recurrencia"], k)
    return out


def _prog_dict(r, hechos, desde, hasta, hoy):
    p = dict(r)
    oc = ocurrencias(p, desde, hasta, hechos)
    p["recurrencia_txt"] = RECURRENCIAS.get(p["recurrencia"] or "", p["recurrencia"])
    p["proximas"] = [f for f in oc if f >= hoy][:6]
    p["vencidas"] = [f for f in oc if f < hoy]
    p["hechos"] = sorted(hechos)
    p["terminado"] = not oc and (not p["recurrencia"] or (p.get("hasta") and p["hasta"] < hoy))
    return p


def listar_programados(dias=120, incluir_terminados=False):
    hoy = _hoy()
    hasta = (date.fromisoformat(hoy) + timedelta(days=dias)).isoformat()
    conn = get_connection()
    try:
        progs = [dict(r) for r in conn.execute("SELECT * FROM teso_programados WHERE estado='activo' ORDER BY fecha, id").fetchall()]
        hechos = {}
        for r in conn.execute("SELECT prog_id, fecha FROM teso_prog_hechos").fetchall():
            hechos.setdefault(r["prog_id"], set()).add(r["fecha"])
        cajas = {r["id"]: r["nombre"] for r in conn.execute("SELECT id, nombre FROM cajas").fetchall()}
    finally:
        conn.close()
    out = []
    for p in progs:
        d = _prog_dict(p, hechos.get(p["id"], set()), "0000-01-01", hasta, hoy)
        if d["terminado"] and not incluir_terminados:
            continue
        d["caja"] = cajas.get(d["caja_id"], "")
        out.append(d)
    return out


def guardar_programado(d, pid=None):
    tipo = d.get("tipo") if d.get("tipo") in ("egreso", "ingreso") else "egreso"
    monto = round(_num(d.get("monto")))
    fecha = _fecha(d.get("fecha"))
    if monto <= 0:
        return {"ok": False, "msg": "Poné el monto"}
    if not fecha:
        return {"ok": False, "msg": "Poné la fecha (o el primer vencimiento)"}
    desc = _txt(d.get("descripcion"), 200)
    if not desc and not _txt(d.get("concepto"), 120):
        return {"ok": False, "msg": "Escribí qué es (ej: sueldo de taller, cuota del camión)"}
    rec = d.get("recurrencia") if d.get("recurrencia") in RECURRENCIAS else ""
    cfg = config()
    necesita = tipo == "egreso" and monto >= cfg["umbral_aprobacion"] > 0
    valores = {
        "tipo": tipo, "caja_id": int(_num(d.get("caja_id"))), "concepto_id": int(_num(d.get("concepto_id"))),
        "concepto": _txt(d.get("concepto"), 120), "descripcion": desc, "contraparte": _txt(d.get("contraparte"), 150),
        "ruc": _txt(d.get("ruc"), 20), "doc_numero": _txt(d.get("doc_numero"), 30), "monto": monto, "fecha": fecha,
        "recurrencia": rec, "hasta": _fecha(d.get("hasta")) if rec else "",
        "medio": d.get("medio") if d.get("medio") in ("efectivo", "transferencia", "cheque", "tarjeta", "vale") else "transferencia",
        "prioridad": d.get("prioridad") if d.get("prioridad") in ("alta", "normal", "baja") else "normal",
        "nota": _txt(d.get("nota"), 500),
    }
    conn = get_connection()
    try:
        if pid:
            act = conn.execute("SELECT * FROM teso_programados WHERE id=?", (pid,)).fetchone()
            if not act:
                return {"ok": False, "msg": "Ya no existe"}
            # si cambió el monto de un pago grande, vuelve a pedir aprobación
            if necesita and (act["aprobacion"] != "aprobado" or float(act["monto"]) < monto):
                valores["aprobacion"] = "admin" if session.get("rol") == "admin" else "pendiente"
            elif not necesita:
                valores["aprobacion"] = "no_requiere"
            if valores.get("aprobacion") == "admin":
                valores.update(aprobacion="aprobado", aprobado_por=_quien(), aprobado_el=ahora_iso())
            sets = ", ".join(f"{k}=?" for k in valores)
            conn.execute(f"UPDATE teso_programados SET {sets} WHERE id=?", list(valores.values()) + [pid])
        else:
            if necesita:
                if session.get("rol") == "admin":
                    valores.update(aprobacion="aprobado", aprobado_por=_quien(), aprobado_el=ahora_iso())
                else:
                    valores["aprobacion"] = "pendiente"
            else:
                valores["aprobacion"] = "no_requiere"
            valores.update(creado_por=_quien(), creado_el=ahora_iso())
            cols = list(valores)
            cur = conn.execute(f"INSERT INTO teso_programados ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})",
                               [valores[c] for c in cols])
            pid = cur.lastrowid
        conn.commit()
    except Exception as e:
        _rollback(conn)
        return {"ok": False, "msg": f"No se pudo guardar: {e}"}
    finally:
        conn.close()
    msg = "Guardado"
    if valores.get("aprobacion") == "pendiente":
        msg += f" · queda esperando la aprobación de administración (es de {_gs(cfg['umbral_aprobacion'])} o más)"
    return {"ok": True, "id": pid, "msg": msg}


def _gs(v):
    return "Gs. " + f"{round(v or 0):,}".replace(",", ".")


def aprobar(pid, aprobar_si, motivo=""):
    conn = get_connection()
    try:
        p = conn.execute("SELECT * FROM teso_programados WHERE id=?", (pid,)).fetchone()
        if not p:
            return {"ok": False, "msg": "Ya no existe"}
        if aprobar_si:
            conn.execute("UPDATE teso_programados SET aprobacion='aprobado', aprobado_por=?, aprobado_el=? WHERE id=?",
                         (_quien(), ahora_iso(), pid))
        else:
            conn.execute("UPDATE teso_programados SET aprobacion='rechazado', aprobado_por=?, aprobado_el=?, nota=? WHERE id=?",
                         (_quien(), ahora_iso(), _txt(f"Rechazado: {motivo}. {p['nota'] or ''}", 500), pid))
        conn.commit()
    finally:
        conn.close()
    return {"ok": True, "msg": "Aprobado: ya se puede pagar" if aprobar_si else "Rechazado"}


def cumplir(pid, fecha, d):
    """Pagar (o cobrar) una vez un programado: se carga el movimiento en la
    caja elegida y esa fecha queda hecha. 'saltar' la marca sin movimiento."""
    import cajas
    conn = get_connection()
    try:
        p = conn.execute("SELECT * FROM teso_programados WHERE id=?", (pid,)).fetchone()
        if not p:
            return {"ok": False, "msg": "Ya no existe"}
        p = dict(p)
        if conn.execute("SELECT 1 FROM teso_prog_hechos WHERE prog_id=? AND fecha=?", (pid, fecha)).fetchone():
            return {"ok": False, "msg": "Ese vencimiento ya estaba hecho"}
    finally:
        conn.close()
    saltar = d.get("accion") == "saltar"
    if not saltar and p["tipo"] == "egreso" and p["aprobacion"] in ("pendiente", "rechazado"):
        return {"ok": False, "msg": "Este pago necesita la aprobación de administración antes de pagarse"}
    mov_id, monto = 0, round(_num(d.get("monto")) or p["monto"])
    if not saltar:
        caja_id = int(_num(d.get("caja_id")) or p["caja_id"])
        if not caja_id:
            return {"ok": False, "msg": "Elegí de qué caja o cuenta sale (o a cuál entra)"}
        mov = {
            "tipo": p["tipo"], "fecha_oper": _fecha(d.get("fecha_pago")) or _hoy(), "fecha_emision": _fecha(d.get("fecha_pago")) or _hoy(),
            "contraparte": p["contraparte"], "ruc": p["ruc"],
            "doc_tipo": d.get("doc_tipo") or ("recibo" if p["tipo"] == "ingreso" else "interno"),
            "doc_numero": d.get("doc_numero") or p["doc_numero"], "timbrado": d.get("timbrado") or "",
            "medio": d.get("medio") or p["medio"], "medio_detalle": d.get("medio_detalle") or "",
            "observacion": f"{p['descripcion'] or p['concepto']} · vencimiento del {fecha[8:10]}/{fecha[5:7]}/{fecha[:4]}",
            "no_imputa": d.get("doc_tipo") in (None, "", "interno", "recibo"),
            "detalle": [{"concepto_id": p["concepto_id"], "concepto": p["concepto"] or p["descripcion"] or "Pago programado",
                         "detalle": p["descripcion"], "tipo_iva": d.get("tipo_iva") or "exento", "total": monto}],
            "confirmar": True,
        }
        r = cajas.crear_mov(caja_id, mov)
        if not r.get("ok"):
            return r
        mov_id = r["id"]
    conn = get_connection()
    try:
        conn.execute("""INSERT INTO teso_prog_hechos (prog_id, fecha, accion, mov_id, monto, hecho_por, hecho_el)
                        VALUES (?,?,?,?,?,?,?)""", (pid, fecha, "saltado" if saltar else "pagado", mov_id, monto, _quien(), ahora_iso()))
        conn.commit()
    finally:
        conn.close()
    if saltar:
        return {"ok": True, "msg": "Ese vencimiento quedó marcado como que no se paga"}
    return {"ok": True, "mov_id": mov_id, "msg": ("Cobrado" if p["tipo"] == "ingreso" else "Pagado") + " y cargado en la caja"}


def borrar_programado(pid):
    conn = get_connection()
    try:
        conn.execute("UPDATE teso_programados SET estado='borrado' WHERE id=?", (pid,))
        conn.commit()
    finally:
        conn.close()
    return {"ok": True, "msg": "Borrado (lo que ya se pagó queda en las cajas)"}


# ════════════════════════════════════════════════════════════════════════════
# CHEQUES
# ════════════════════════════════════════════════════════════════════════════

def _cheque_dict(r, hoy):
    c = dict(r)
    c["estado_txt"] = ESTADOS_CHEQUE.get(c["sentido"], {}).get(c["estado"], c["estado"])
    try:
        c["historial"] = json.loads(c.get("historial") or "[]")
    except ValueError:
        c["historial"] = []
    fp = date.fromisoformat(c["fecha_pago"])
    h = date.fromisoformat(hoy)
    c["dias"] = (fp - h).days
    avisos = []
    if c["sentido"] == "recibido" and c["estado"] == "cartera":
        limite = fp + timedelta(days=CHEQUE_PRESENTACION_DIAS)
        c["presentar_hasta"] = limite.isoformat()
        if h < fp:
            avisos.append(("info", f"Se puede depositar desde el {fp.strftime('%d/%m')}" + (" (si se deposita antes, el banco lo rechaza)" if c["tipo"] == "diferido" else "")))
        elif h <= limite:
            dl = (limite - h).days
            avisos.append(("mal" if dl <= 5 else "ojo", f"Depositalo: quedan {dl} día{'s' if dl != 1 else ''} para presentarlo (hasta el {limite.strftime('%d/%m')})"))
        else:
            avisos.append(("mal", f"Venció el plazo de presentación el {limite.strftime('%d/%m')}"))
    if c["sentido"] == "emitido" and c["estado"] in ("emitido", "entregado"):
        if 0 <= c["dias"] <= 3:
            avisos.append(("ojo", f"Se debita {'hoy' if c['dias'] == 0 else 'en ' + str(c['dias']) + ' días'}: que haya fondos en la cuenta"))
        elif c["dias"] < 0:
            avisos.append(("ojo", "Ya pasó la fecha de pago: marcalo como debitado cuando salga en el extracto"))
    c["avisos"] = [{"nivel": a, "msg": m} for a, m in avisos]
    return c


def listar_cheques(sentido="", estado="", q=""):
    hoy = _hoy()
    conn = get_connection()
    try:
        sql = "SELECT * FROM teso_cheques WHERE 1=1"
        p = []
        if sentido in ("recibido", "emitido"):
            sql += " AND sentido=?"
            p.append(sentido)
        if estado == "abiertos":
            sql += " AND estado IN ('cartera','depositado','emitido','entregado')"
        elif estado:
            sql += " AND estado=?"
            p.append(estado)
        if q:
            like = f"%{q.strip().lower()}%"
            sql += " AND (LOWER(persona) LIKE ? OR LOWER(numero) LIKE ? OR LOWER(banco) LIKE ? OR LOWER(concepto) LIKE ?)"
            p += [like, like, like, like]
        sql += " ORDER BY fecha_pago, id LIMIT 600"
        filas = [_cheque_dict(r, hoy) for r in conn.execute(sql, p).fetchall()]
        cajas = {r["id"]: r["nombre"] for r in conn.execute("SELECT id, nombre FROM cajas").fetchall()}
    finally:
        conn.close()
    for f in filas:
        f["caja"] = cajas.get(f["caja_id"], "")
    abiertos = [f for f in filas if f["estado"] in ("cartera", "depositado", "emitido", "entregado")]
    tot = {
        "cartera": sum(f["monto"] for f in abiertos if f["sentido"] == "recibido"),
        "n_cartera": sum(1 for f in abiertos if f["sentido"] == "recibido"),
        "por_debitar": sum(f["monto"] for f in abiertos if f["sentido"] == "emitido"),
        "n_por_debitar": sum(1 for f in abiertos if f["sentido"] == "emitido"),
        "avisos": sum(1 for f in abiertos if any(a["nivel"] in ("mal", "ojo") for a in f["avisos"])),
    }
    return {"cheques": filas, "totales": tot}


def _hist(c, txt):
    h = c.get("historial") if isinstance(c.get("historial"), list) else json.loads(c.get("historial") or "[]")
    h.append({"el": ahora_iso()[:16], "por": _quien(), "txt": txt})
    return json.dumps(h, ensure_ascii=False)


def guardar_cheque(d, cid=None):
    sentido = d.get("sentido") if d.get("sentido") in ("recibido", "emitido") else "recibido"
    tipo = "comun" if d.get("tipo") == "comun" else "diferido"
    monto = round(_num(d.get("monto")))
    fe = _fecha(d.get("fecha_emision")) or _hoy()
    fp = _fecha(d.get("fecha_pago")) or fe
    if monto <= 0:
        return {"ok": False, "msg": "Poné el monto"}
    if not _txt(d.get("numero"), 30):
        return {"ok": False, "msg": "Poné el número del cheque"}
    if fp < fe:
        return {"ok": False, "msg": "La fecha de pago no puede ser anterior a la de emisión"}
    dias = (date.fromisoformat(fp) - date.fromisoformat(fe)).days
    if tipo == "diferido" and dias > CHEQUE_DIFERIDO_MAX_DIAS:
        return {"ok": False, "msg": f"Un cheque de pago diferido no puede ser a más de {CHEQUE_DIFERIDO_MAX_DIAS} días de la emisión (Ley 805/96): este es a {dias}"}
    if tipo == "comun":
        fp = fe
    v = {"sentido": sentido, "tipo": tipo, "banco": _txt(d.get("banco"), 60), "numero": _txt(d.get("numero"), 30),
         "persona": _txt(d.get("persona"), 150), "ruc": _txt(d.get("ruc"), 20), "monto": monto,
         "fecha_emision": fe, "fecha_pago": fp, "caja_id": int(_num(d.get("caja_id"))),
         "concepto": _txt(d.get("concepto"), 150), "nota": _txt(d.get("nota"), 500)}
    if sentido == "emitido" and not v["caja_id"]:
        return {"ok": False, "msg": "Elegí la cuenta del banco de la que sale el cheque"}
    conn = get_connection()
    try:
        dup = conn.execute("""SELECT id FROM teso_cheques WHERE sentido=? AND LOWER(banco)=? AND numero=? AND id<>?""",
                           (sentido, v["banco"].lower(), v["numero"], cid or 0)).fetchone()
        if dup:
            return {"ok": False, "msg": "Ese cheque (mismo banco y número) ya está cargado"}
        if cid:
            act = conn.execute("SELECT * FROM teso_cheques WHERE id=?", (cid,)).fetchone()
            if not act:
                return {"ok": False, "msg": "Ya no existe"}
            if act["mov_id"]:
                return {"ok": False, "msg": "Ya está cobrado o debitado: no se puede cambiar"}
            v["historial"] = _hist(dict(act), "Corregido")
            conn.execute(f"UPDATE teso_cheques SET {', '.join(k + '=?' for k in v)} WHERE id=?", list(v.values()) + [cid])
        else:
            v.update(estado="cartera" if sentido == "recibido" else "emitido", creado_por=_quien(), creado_el=ahora_iso())
            v["historial"] = json.dumps([{"el": ahora_iso()[:16], "por": _quien(), "txt": "Cargado"}], ensure_ascii=False)
            cols = list(v)
            cur = conn.execute(f"INSERT INTO teso_cheques ({','.join(cols)}) VALUES ({','.join('?' * len(cols))})", [v[c] for c in cols])
            cid = cur.lastrowid
        conn.commit()
    except Exception as e:
        _rollback(conn)
        return {"ok": False, "msg": f"No se pudo guardar: {e}"}
    finally:
        conn.close()
    return {"ok": True, "id": cid, "msg": "Cheque guardado"}


def accion_cheque(cid, accion, d):
    """depositar · cobrar (acreditado) · endosar · rechazar · entregar ·
    debitar · anular. Cobrar y debitar cargan el movimiento en la cuenta."""
    import cajas
    conn = get_connection()
    try:
        c = conn.execute("SELECT * FROM teso_cheques WHERE id=?", (cid,)).fetchone()
    finally:
        conn.close()
    if not c:
        return {"ok": False, "msg": "Ya no existe"}
    c = dict(c)
    hoy = _hoy()
    nuevo, txt, mov = None, "", None
    s, e = c["sentido"], c["estado"]
    if s == "recibido":
        if accion == "depositar" and e == "cartera":
            if c["tipo"] == "diferido" and hoy < c["fecha_pago"]:
                return {"ok": False, "msg": f"Es diferido: recién se puede depositar desde el {c['fecha_pago'][8:10]}/{c['fecha_pago'][5:7]}"}
            caja_id = int(_num(d.get("caja_id")))
            if not caja_id:
                return {"ok": False, "msg": "Elegí en qué cuenta se deposita"}
            c["caja_id"] = caja_id
            nuevo, txt = "depositado", "Depositado"
        elif accion == "cobrar" and e in ("cartera", "depositado"):
            caja_id = int(_num(d.get("caja_id")) or c["caja_id"])
            if not caja_id:
                return {"ok": False, "msg": "Elegí en qué cuenta se acreditó"}
            c["caja_id"] = caja_id
            nuevo, txt = "cobrado", "Acreditado en la cuenta"
            mov = {"tipo": "ingreso", "medio": "cheque"}
        elif accion == "endosar" and e == "cartera":
            a_quien = _txt(d.get("a_quien"), 150)
            if not a_quien:
                return {"ok": False, "msg": "¿A quién se le endosó?"}
            nuevo, txt = "endosado", f"Endosado a {a_quien}"
        elif accion == "rechazar" and e in ("cartera", "depositado", "cobrado"):
            nuevo, txt = "rechazado", f"Rechazado por el banco: {_txt(d.get('motivo'), 150) or 'sin motivo'}"
            if e == "cobrado" and c["mov_id"]:
                mov = {"tipo": "egreso", "medio": "cheque", "reverso": True}
        else:
            return {"ok": False, "msg": "No se puede hacer eso con este cheque"}
    else:
        if accion == "entregar" and e == "emitido":
            nuevo, txt = "entregado", f"Entregado a {c['persona'] or 'el beneficiario'}"
        elif accion == "debitar" and e in ("emitido", "entregado"):
            nuevo, txt = "debitado", "Debitado de la cuenta"
            mov = {"tipo": "egreso", "medio": "cheque"}
        elif accion == "anular" and e in ("emitido", "entregado"):
            nuevo, txt = "anulado", f"Anulado: {_txt(d.get('motivo'), 150) or 'sin motivo'}"
        elif accion == "rechazar" and e in ("entregado", "debitado"):
            nuevo, txt = "rechazado", f"Rechazado: {_txt(d.get('motivo'), 150) or 'sin motivo'}"
            if e == "debitado" and c["mov_id"]:
                mov = {"tipo": "ingreso", "medio": "cheque", "reverso": True}
        else:
            return {"ok": False, "msg": "No se puede hacer eso con este cheque"}
    mov_id = c["mov_id"]
    if mov:
        rev = mov.get("reverso")
        r = cajas.crear_mov(c["caja_id"], {
            "tipo": mov["tipo"], "fecha_oper": _fecha(d.get("fecha")) or hoy, "fecha_emision": _fecha(d.get("fecha")) or hoy,
            "contraparte": c["persona"], "ruc": c["ruc"], "doc_tipo": "interno",
            "doc_numero": f"CH {c['numero']}", "medio": "cheque", "medio_detalle": f"{c['banco']} N° {c['numero']}",
            "observacion": (("Reverso por cheque rechazado. " if rev else "") + (c["concepto"] or "")).strip(),
            "no_imputa": True, "confirmar": True,
            "detalle": [{"concepto": ("Cheque rechazado" if rev else ("Cobro de cheque" if mov["tipo"] == "ingreso" else "Pago con cheque")),
                         "detalle": f"{c['banco']} N° {c['numero']} · {c['persona']}", "tipo_iva": "exento", "total": c["monto"]}],
        })
        if not r.get("ok"):
            return r
        mov_id = r["id"]
    conn = get_connection()
    try:
        conn.execute("UPDATE teso_cheques SET estado=?, caja_id=?, mov_id=?, historial=? WHERE id=?",
                     (nuevo, c["caja_id"], mov_id, _hist(c, txt), cid))
        conn.commit()
    finally:
        conn.close()
    return {"ok": True, "msg": txt + (" · movimiento cargado en la cuenta" if mov else "")}


# ════════════════════════════════════════════════════════════════════════════
# FLUJO DE CAJA PROYECTADO
# ════════════════════════════════════════════════════════════════════════════

def _lunes(d):
    return d - timedelta(days=d.weekday())


def flujo(semanas=13, cajas_ids=None, habitual=True):
    """Semana por semana: saldo inicial, lo que entra, lo que sale y el saldo
    final. Lo vencido sin pagar se cuenta en la semana actual."""
    import cajas as cj
    hoy = date.fromisoformat(_hoy())
    semanas = max(4, min(26, int(semanas or 13)))
    ini = _lunes(hoy)
    fin = ini + timedelta(days=7 * semanas - 1)
    todas = cj.listar_cajas()
    sel = [c for c in todas if not cajas_ids or c["id"] in cajas_ids]
    saldo0 = sum(float(c["saldo"] or 0) for c in sel)
    minimo = sum(float(c.get("saldo_minimo") or 0) for c in sel)
    ids_sel = {c["id"] for c in sel}
    filas = [{"desde": (ini + timedelta(days=7 * k)).isoformat(), "hasta": (ini + timedelta(days=7 * k + 6)).isoformat(),
              "ingresos": 0.0, "egresos": 0.0, "habitual": 0.0, "items": []} for k in range(semanas)]

    def poner(fecha, tipo, monto, txt, origen, ref=None, vencido=False):
        f = max(date.fromisoformat(fecha), hoy)
        k = (f - ini).days // 7
        if k < 0 or k >= semanas:
            return
        filas[k]["ingresos" if tipo == "ingreso" else "egresos"] += monto
        filas[k]["items"].append({"fecha": fecha, "tipo": tipo, "monto": monto, "txt": txt, "origen": origen,
                                  "ref": ref, "vencido": vencido})

    hasta = fin.isoformat()
    for p in listar_programados(dias=7 * semanas + 7):
        if p["caja_id"] and cajas_ids and p["caja_id"] not in ids_sel:
            continue
        if p["aprobacion"] == "rechazado":
            continue
        for f in p["vencidas"] + [x for x in ocurrencias(p, "", hasta, set(p["hechos"])) if x >= hoy.isoformat()]:
            if f > hasta:
                continue
            poner(f, p["tipo"], p["monto"], p["descripcion"] or p["concepto"], "programado", {"id": p["id"], "fecha": f},
                  vencido=f < hoy.isoformat())
    ch = listar_cheques(estado="abiertos")["cheques"]
    for c in ch:
        if cajas_ids and c["caja_id"] and c["caja_id"] not in ids_sel:
            continue
        tipo = "ingreso" if c["sentido"] == "recibido" else "egreso"
        poner(c["fecha_pago"], tipo, c["monto"], f"Cheque {c['banco']} N° {c['numero']} · {c['persona']}", "cheque",
              {"id": c["id"]}, vencido=c["fecha_pago"] < hoy.isoformat())
    # el gasto habitual: lo que sale cada semana sin estar programado (promedio de 8 semanas)
    prom = 0.0
    if habitual:
        desde8 = (ini - timedelta(days=56)).isoformat()
        conn = get_connection()
        try:
            prog_movs = {r["mov_id"] for r in conn.execute("SELECT mov_id FROM teso_prog_hechos WHERE mov_id>0").fetchall()}
            ch_movs = {r["mov_id"] for r in conn.execute("SELECT mov_id FROM teso_cheques WHERE mov_id>0").fetchall()}
            q = """SELECT id, caja_id, total_gs FROM caja_movs WHERE estado='vigente' AND tipo='egreso'
                   AND COALESCE(transf_id,0)=0 AND fecha_oper>=? AND fecha_oper<?"""
            tot = 0.0
            for r in conn.execute(q, (desde8, ini.isoformat())).fetchall():
                if r["caja_id"] in ids_sel and r["id"] not in prog_movs and r["id"] not in ch_movs:
                    tot += float(r["total_gs"] or 0)
        finally:
            conn.close()
        prom = round(tot / 8)
        if prom:
            # de esta semana se descuenta lo que ya salió
            conn = get_connection()
            try:
                ya = conn.execute("""SELECT COALESCE(SUM(total_gs),0) AS t FROM caja_movs WHERE estado='vigente' AND tipo='egreso'
                                     AND COALESCE(transf_id,0)=0 AND fecha_oper>=? AND fecha_oper<=?""",
                                  (ini.isoformat(), hoy.isoformat())).fetchone()["t"] or 0
            finally:
                conn.close()
            for k, fl in enumerate(filas):
                fl["habitual"] = max(0.0, prom - float(ya)) if k == 0 else prom
                fl["egresos"] += fl["habitual"]
    saldo = saldo0
    alerta = None
    for fl in filas:
        fl["inicial"] = round(saldo)
        saldo += fl["ingresos"] - fl["egresos"]
        fl["final"] = round(saldo)
        fl["ingresos"], fl["egresos"] = round(fl["ingresos"]), round(fl["egresos"])
        fl["items"].sort(key=lambda x: (x["fecha"], x["tipo"]))
        if alerta is None and (saldo < 0 or (minimo and saldo < minimo)):
            alerta = {"desde": fl["desde"], "saldo": round(saldo), "negativo": saldo < 0}
    vencidos = [i for fl in filas for i in fl["items"] if i["vencido"]]
    return {"hoy": hoy.isoformat(), "saldo_actual": round(saldo0), "minimo": round(minimo), "semanas": filas,
            "habitual_semanal": prom, "alerta": alerta, "minimo_proyectado": min([f["final"] for f in filas] + [round(saldo0)]),
            "vencidos": {"n": len(vencidos), "egresos": round(sum(i["monto"] for i in vencidos if i["tipo"] == "egreso")),
                         "ingresos": round(sum(i["monto"] for i in vencidos if i["tipo"] == "ingreso"))},
            "cajas": [{"id": c["id"], "nombre": c["nombre"], "tipo": c.get("tipo") or "caja", "saldo": round(c["saldo"]),
                       "banco": c.get("banco") or "", "numero_cuenta": c.get("numero_cuenta") or "",
                       "saldo_minimo": c.get("saldo_minimo") or 0, "sel": c["id"] in ids_sel} for c in todas]}


# ════════════════════════════════════════════════════════════════════════════
# CONCILIACIÓN BANCARIA
# ════════════════════════════════════════════════════════════════════════════

ENC = {
    "fecha": ["fecha", "fecha operacion", "fecha valor", "fec", "date", "f. operacion", "fecha mov"],
    "desc": ["descripcion", "concepto", "detalle", "movimiento", "glosa", "description", "transaccion", "operacion"],
    "ref": ["referencia", "comprobante", "nro", "documento", "ref", "numero", "n° comprobante", "nro. comprobante"],
    "debito": ["debito", "debe", "cargo", "egreso", "retiro", "debitos", "salida"],
    "credito": ["credito", "haber", "abono", "ingreso", "deposito", "creditos", "entrada"],
    "monto": ["importe", "monto", "valor", "amount"],
    "saldo": ["saldo", "balance", "saldo disponible"],
}


def _norm(t):
    import unicodedata
    t = unicodedata.normalize("NFD", str(t or "").lower())
    return re.sub(r"\s+", " ", "".join(c for c in t if unicodedata.category(c) != "Mn")).strip(" .:")


def _leer_tabla(nombre, crudo):
    """Filas (listas) de un Excel o CSV del banco."""
    if nombre.lower().endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(crudo), read_only=True, data_only=True)
        ws = wb.active
        return [list(r) for r in ws.iter_rows(values_only=True)]
    if nombre.lower().endswith(".xls"):
        raise ValueError("Guardá el extracto como .xlsx o .csv (el formato .xls viejo no se puede leer)")
    txt = None
    for enc in ("utf-8-sig", "latin-1"):
        try:
            txt = crudo.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if txt is None:
        raise ValueError("No se pudo leer el texto del archivo")
    # el separador: el que más se repite en las primeras líneas (; , tab o |)
    muestra = txt[:6000]
    sep = max(";,\t|", key=lambda ch: muestra.count(ch))
    return [r for r in csv.reader(io.StringIO(txt), delimiter=sep)]


def _num_banco(v):
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("Gs", "").replace("₲", "").replace(" ", "")
    neg = s.startswith("-") or (s.startswith("(") and s.endswith(")")) or s.endswith("-")
    s = s.strip("-()")
    if re.fullmatch(r"\d{1,3}(\.\d{3})+(,\d+)?", s):
        s = s.replace(".", "").replace(",", ".")
    elif re.fullmatch(r"\d{1,3}(,\d{3})+(\.\d+)?", s):
        s = s.replace(",", "")
    else:
        s = s.replace(",", ".")
    try:
        n = float(s)
    except ValueError:
        return None
    return -n if neg else n


def leer_extracto(nombre, crudo):
    filas = _leer_tabla(nombre, crudo)
    # buscar la fila de encabezados
    cab_i, mapa = None, {}
    for i, f in enumerate(filas[:40]):
        m = {}
        for j, celda in enumerate(f):
            t = _norm(celda)
            if not t:
                continue
            for k, nombres in ENC.items():
                if k not in m and any(t == n or t.startswith(n + " ") for n in nombres):
                    m[k] = j
        if "fecha" in m and ("monto" in m or "debito" in m or "credito" in m):
            cab_i, mapa = i, m
            break
    if cab_i is None:
        raise ValueError("No encontré las columnas: el extracto tiene que tener Fecha y Débito/Crédito (o Importe)")
    out = []
    for f in filas[cab_i + 1:]:
        g = lambda k: f[mapa[k]] if k in mapa and mapa[k] < len(f) else None
        fv = g("fecha")
        fecha = fv.date().isoformat() if isinstance(fv, datetime) else (fv.isoformat() if isinstance(fv, date) else _fecha(fv))
        if not fecha:
            continue
        if "debito" in mapa or "credito" in mapa:
            deb, cre = _num_banco(g("debito")) or 0, _num_banco(g("credito")) or 0
            monto = abs(cre) - abs(deb)
        else:
            monto = _num_banco(g("monto")) or 0
        if not monto:
            continue
        out.append({"fecha": fecha, "descripcion": _txt(g("desc"), 250), "referencia": _txt(g("ref"), 60),
                    "monto": round(monto, 2), "saldo": _num_banco(g("saldo"))})
    if not out:
        raise ValueError("No encontré movimientos en el archivo")
    return out


def importar_extracto(caja_id, nombre, crudo):
    try:
        lineas = leer_extracto(nombre, crudo)
    except ValueError as e:
        return {"ok": False, "msg": str(e)}
    except Exception as e:
        return {"ok": False, "msg": f"No se pudo leer el archivo: {e}"}
    lote = ahora_iso()[:16]
    nuevas = repetidas = 0
    conn = get_connection()
    try:
        for ln in lineas:
            ya = conn.execute("""SELECT 1 FROM teso_extracto WHERE caja_id=? AND fecha=? AND monto=? AND descripcion=? AND referencia=?""",
                              (caja_id, ln["fecha"], ln["monto"], ln["descripcion"], ln["referencia"])).fetchone()
            if ya:
                repetidas += 1
                continue
            conn.execute("""INSERT INTO teso_extracto (caja_id, fecha, descripcion, referencia, monto, saldo, lote, importado_por, importado_el)
                            VALUES (?,?,?,?,?,?,?,?,?)""",
                         (caja_id, ln["fecha"], ln["descripcion"], ln["referencia"], ln["monto"], ln["saldo"], lote, _quien(), ahora_iso()))
            nuevas += 1
        conn.commit()
    finally:
        conn.close()
    auto = conciliar_auto(caja_id)
    return {"ok": True, "nuevas": nuevas, "repetidas": repetidas, "conciliadas": auto,
            "msg": f"{nuevas} líneas nuevas" + (f" ({repetidas} ya estaban)" if repetidas else "") + f" · {auto} conciliadas solas"}


def conciliar_auto(caja_id, tolerancia_dias=5):
    """Cruza cada línea del extracto con un movimiento de la cuenta: mismo
    monto y signo, fecha a menos de 5 días; si hay varios, el más cercano."""
    conn = get_connection()
    try:
        ext = [dict(r) for r in conn.execute("""SELECT * FROM teso_extracto WHERE caja_id=? AND estado='pendiente'
                                               ORDER BY fecha, id""", (caja_id,)).fetchall()]
        movs = [dict(r) for r in conn.execute("""SELECT id, tipo, fecha_oper, total_gs, doc_numero, medio_detalle, contraparte
                                                FROM caja_movs WHERE caja_id=? AND estado='vigente' AND COALESCE(extracto_id,0)=0""",
                                             (caja_id,)).fetchall()]
        n = 0
        usados = set()
        for e in ext:
            signo = 1 if e["monto"] > 0 else -1
            fe = date.fromisoformat(e["fecha"])
            cand = []
            for m in movs:
                if m["id"] in usados:
                    continue
                ms = 1 if m["tipo"] == "ingreso" else -1
                if ms != signo or abs(abs(e["monto"]) - float(m["total_gs"] or 0)) > 0.5:
                    continue
                dd = abs((date.fromisoformat(m["fecha_oper"]) - fe).days)
                if dd > tolerancia_dias:
                    continue
                # la referencia (número de cheque o transferencia) ayuda a desempatar
                ref = re.sub(r"\D", "", e["referencia"] or "")
                bonus = -3 if ref and len(ref) >= 3 and ref in re.sub(r"\D", "", f"{m['doc_numero']} {m['medio_detalle']}") else 0
                cand.append((dd + bonus, m["id"]))
            if not cand:
                continue
            cand.sort()
            if len(cand) > 1 and cand[0][0] == cand[1][0]:
                continue        # dos iguales: mejor que lo decida una persona
            mid = cand[0][1]
            usados.add(mid)
            conn.execute("UPDATE teso_extracto SET mov_id=?, estado='conciliado', como='auto' WHERE id=?", (mid, e["id"]))
            conn.execute("UPDATE caja_movs SET extracto_id=? WHERE id=?", (e["id"], mid))
            n += 1
        conn.commit()
        return n
    finally:
        conn.close()


def estado_conciliacion(caja_id, desde="", hasta=""):
    conn = get_connection()
    try:
        caja = conn.execute("SELECT * FROM cajas WHERE id=?", (caja_id,)).fetchone()
        if not caja:
            return None
        q = "SELECT * FROM teso_extracto WHERE caja_id=?"
        p = [caja_id]
        if desde:
            q += " AND fecha>=?"
            p.append(desde)
        if hasta:
            q += " AND fecha<=?"
            p.append(hasta)
        ext = [dict(r) for r in conn.execute(q + " ORDER BY fecha DESC, id DESC LIMIT 1500", p).fetchall()]
        q2 = """SELECT id, numero, tipo, fecha_oper, contraparte, total_gs, medio, doc_numero, medio_detalle, observacion,
                       COALESCE(extracto_id,0) AS extracto_id FROM caja_movs WHERE caja_id=? AND estado='vigente'"""
        p2 = [caja_id]
        primera = min([e["fecha"] for e in ext], default="")
        if desde or primera:
            q2 += " AND fecha_oper>=?"
            p2.append(desde or (date.fromisoformat(primera) - timedelta(days=10)).isoformat())
        movs = [dict(r) for r in conn.execute(q2 + " ORDER BY fecha_oper DESC, numero DESC LIMIT 1500", p2).fetchall()]
        conc = {}
        mids = [m["id"] for m in movs]
        if mids:
            for r in conn.execute(f"SELECT mov_id, concepto FROM caja_mov_det WHERE mov_id IN ({','.join('?' * len(mids))}) ORDER BY orden",
                                  mids).fetchall():
                conc.setdefault(r["mov_id"], r["concepto"])
        for m in movs:
            m["concepto"] = conc.get(m["id"], "")
    finally:
        conn.close()
    ult = next((e for e in ext if e["saldo"] is not None), None)
    import cajas as cj
    saldo_sistema = next((c["saldo"] for c in cj.listar_cajas(True) if c["id"] == caja_id), 0)
    return {"caja": dict(caja), "extracto": ext, "movs": movs,
            "pendientes_banco": sum(1 for e in ext if e["estado"] == "pendiente"),
            "pendientes_sistema": sum(1 for m in movs if not m["extracto_id"]),
            "saldo_banco": ult["saldo"] if ult else None, "saldo_banco_fecha": ult["fecha"] if ult else "",
            "saldo_sistema": saldo_sistema}


def conciliar_manual(eid, mid):
    conn = get_connection()
    try:
        e = conn.execute("SELECT * FROM teso_extracto WHERE id=?", (eid,)).fetchone()
        m = conn.execute("SELECT id, caja_id, extracto_id FROM caja_movs WHERE id=?", (mid,)).fetchone()
        if not e or not m or m["caja_id"] != e["caja_id"]:
            return {"ok": False, "msg": "No coinciden la cuenta y el movimiento"}
        if m["extracto_id"]:
            return {"ok": False, "msg": "Ese movimiento ya está conciliado con otra línea"}
        conn.execute("UPDATE teso_extracto SET mov_id=?, estado='conciliado', como='manual' WHERE id=?", (mid, eid))
        conn.execute("UPDATE caja_movs SET extracto_id=? WHERE id=?", (eid, mid))
        conn.commit()
    finally:
        conn.close()
    return {"ok": True, "msg": "Conciliado"}


def desconciliar(eid):
    conn = get_connection()
    try:
        e = conn.execute("SELECT * FROM teso_extracto WHERE id=?", (eid,)).fetchone()
        if not e:
            return {"ok": False, "msg": "Ya no existe"}
        if e["mov_id"]:
            conn.execute("UPDATE caja_movs SET extracto_id=0 WHERE id=?", (e["mov_id"],))
        conn.execute("UPDATE teso_extracto SET mov_id=0, estado='pendiente', como='' WHERE id=?", (eid,))
        conn.commit()
    finally:
        conn.close()
    return {"ok": True, "msg": "Vuelve a pendiente"}


def ignorar(eid):
    conn = get_connection()
    try:
        conn.execute("UPDATE teso_extracto SET estado='ignorado', como='ignorado' WHERE id=? AND estado='pendiente'", (eid,))
        conn.commit()
    finally:
        conn.close()
    return {"ok": True, "msg": "Línea ignorada"}


def crear_desde_linea(eid, d):
    """Lo que solo está en el banco (comisión, interés, débito automático):
    se carga en la cuenta y queda conciliado."""
    import cajas
    conn = get_connection()
    try:
        e = conn.execute("SELECT * FROM teso_extracto WHERE id=?", (eid,)).fetchone()
    finally:
        conn.close()
    if not e or e["estado"] != "pendiente":
        return {"ok": False, "msg": "Esa línea ya no está pendiente"}
    concepto = _txt(d.get("concepto"), 120) or ("Gastos bancarios" if e["monto"] < 0 else "Ingreso bancario")
    r = cajas.crear_mov(e["caja_id"], {
        "tipo": "ingreso" if e["monto"] > 0 else "egreso", "fecha_oper": e["fecha"], "fecha_emision": e["fecha"],
        "contraparte": _txt(d.get("contraparte"), 150) or "Banco", "doc_tipo": "interno", "doc_numero": e["referencia"],
        "medio": "transferencia", "observacion": f"Desde el extracto: {e['descripcion']}", "no_imputa": True, "confirmar": True,
        "detalle": [{"concepto_id": int(_num(d.get("concepto_id"))), "concepto": concepto, "detalle": e["descripcion"],
                     "tipo_iva": d.get("tipo_iva") or "exento", "total": abs(e["monto"])}],
    })
    if not r.get("ok"):
        return r
    return conciliar_manual(eid, r["id"]) | {"msg": "Cargado y conciliado"}


# ════════════════════════════════════════════════════════════════════════════
# RESUMEN (lo que la pantalla muestra arriba)
# ════════════════════════════════════════════════════════════════════════════

def resumen():
    hoy = _hoy()
    en7 = (date.fromisoformat(hoy) + timedelta(days=7)).isoformat()
    progs = listar_programados(dias=7)
    pagar7 = cobrar7 = 0
    vencidos = aprobar_n = 0
    for p in progs:
        for f in p["vencidas"] + p["proximas"]:
            if f <= en7:
                if p["tipo"] == "egreso":
                    pagar7 += p["monto"]
                else:
                    cobrar7 += p["monto"]
            if f < hoy:
                vencidos += 1
        if p["aprobacion"] == "pendiente":
            aprobar_n += 1
    ch = listar_cheques(estado="abiertos")
    avisos_ch = ch["totales"]["avisos"]
    conn = get_connection()
    try:
        pend_ext = conn.execute("SELECT COUNT(*) AS n FROM teso_extracto WHERE estado='pendiente'").fetchone()["n"]
    except Exception:
        _rollback(conn)
        pend_ext = 0
    finally:
        conn.close()
    return {"pagar7": round(pagar7), "cobrar7": round(cobrar7), "vencidos": vencidos, "por_aprobar": aprobar_n,
            "cheques_avisos": avisos_ch, "cartera": round(ch["totales"]["cartera"]),
            "por_debitar": round(ch["totales"]["por_debitar"]), "extracto_pendiente": int(pend_ext or 0)}


# ════════════════════════════════════════════════════════════════════════════
# RUTAS
# ════════════════════════════════════════════════════════════════════════════

def _puede(escribir=False):
    import cajas
    return cajas._puede(escribir)


def _no(msg="Sin permiso", code=403):
    return jsonify({"ok": False, "msg": msg, "error": msg}), code


def _resp(r):
    return jsonify(r), (200 if r.get("ok") else 400)


def _auditar(accion, detalle=""):
    try:
        from database import registrar_auditoria
        registrar_auditoria(usuario=_quien(), rol=session.get("rol") or "", accion=accion, categoria="tesoreria",
                            detalle=str(detalle)[:250])
    except Exception:
        pass


@bp_teso.route("/api/teso/resumen")
def api_resumen():
    if not _puede():
        return _no()
    return jsonify({"ok": True, **resumen(), "config": config(), "bancos": BANCOS, "recurrencias": RECURRENCIAS,
                    "estados_cheque": ESTADOS_CHEQUE, "admin": session.get("rol") == "admin"})


@bp_teso.route("/api/teso/flujo")
def api_flujo():
    if not _puede():
        return _no()
    ids = [int(x) for x in (request.args.get("cajas") or "").split(",") if x.strip().isdigit()]
    return jsonify({"ok": True, **flujo(request.args.get("semanas") or 13, ids or None, request.args.get("habitual", "1") == "1")})


@bp_teso.route("/api/teso/programados", methods=["GET"])
def api_programados():
    if not _puede():
        return _no()
    return jsonify({"ok": True, "programados": listar_programados(dias=int(request.args.get("dias") or 120),
                                                                  incluir_terminados=request.args.get("todos") == "1"),
                    "hoy": _hoy()})


@bp_teso.route("/api/teso/programados", methods=["POST"])
@bp_teso.route("/api/teso/programados/<int:pid>", methods=["PUT"])
def api_programado_guardar(pid=None):
    if not _puede(True):
        return _no()
    r = guardar_programado(request.get_json(silent=True) or {}, pid)
    if r.get("ok"):
        _auditar("Programó un pago/cobro" if not pid else "Cambió un pago/cobro programado", r.get("msg"))
    return _resp(r)


@bp_teso.route("/api/teso/programados/<int:pid>", methods=["DELETE"])
def api_programado_borrar(pid):
    if not _puede(True):
        return _no()
    _auditar("Borró un pago/cobro programado", f"#{pid}")
    return _resp(borrar_programado(pid))


@bp_teso.route("/api/teso/programados/<int:pid>/aprobar", methods=["POST"])
def api_programado_aprobar(pid):
    if session.get("rol") != "admin":
        return _no("Solo administración aprueba pagos")
    d = request.get_json(silent=True) or {}
    r = aprobar(pid, d.get("aprobar", True) is not False, d.get("motivo", ""))
    _auditar("Aprobó un pago" if d.get("aprobar", True) is not False else "Rechazó un pago", f"#{pid}")
    return _resp(r)


@bp_teso.route("/api/teso/programados/<int:pid>/cumplir", methods=["POST"])
def api_programado_cumplir(pid):
    if not _puede(True):
        return _no()
    d = request.get_json(silent=True) or {}
    fecha = _fecha(d.get("fecha"))
    if not fecha:
        return _no("Falta el vencimiento", 400)
    r = cumplir(pid, fecha, d)
    if r.get("ok"):
        _auditar("Pagó/cobró un programado", f"#{pid} {fecha}")
    return _resp(r)


@bp_teso.route("/api/teso/cheques", methods=["GET"])
def api_cheques():
    if not _puede():
        return _no()
    a = request.args
    return jsonify({"ok": True, **listar_cheques(a.get("sentido", ""), a.get("estado", ""), a.get("q", "")), "hoy": _hoy()})


@bp_teso.route("/api/teso/cheques", methods=["POST"])
@bp_teso.route("/api/teso/cheques/<int:cid>", methods=["PUT"])
def api_cheque_guardar(cid=None):
    if not _puede(True):
        return _no()
    r = guardar_cheque(request.get_json(silent=True) or {}, cid)
    if r.get("ok"):
        _auditar("Cargó un cheque" if not cid else "Corrigió un cheque", f"#{r.get('id')}")
    return _resp(r)


@bp_teso.route("/api/teso/cheques/<int:cid>/<accion>", methods=["POST"])
def api_cheque_accion(cid, accion):
    if not _puede(True):
        return _no()
    r = accion_cheque(cid, accion, request.get_json(silent=True) or {})
    if r.get("ok"):
        _auditar(f"Cheque: {accion}", f"#{cid} {r.get('msg')}")
    return _resp(r)


@bp_teso.route("/api/teso/conciliacion/<int:caja_id>")
def api_conciliacion(caja_id):
    if not _puede():
        return _no()
    r = estado_conciliacion(caja_id, request.args.get("desde", ""), request.args.get("hasta", ""))
    if r is None:
        return _no("No existe esa cuenta", 404)
    return jsonify({"ok": True, **r})


@bp_teso.route("/api/teso/conciliacion/<int:caja_id>/importar", methods=["POST"])
def api_importar(caja_id):
    if not _puede(True):
        return _no()
    f = request.files.get("archivo")
    if not f:
        return _no("Elegí el archivo del extracto", 400)
    crudo = f.read()
    if len(crudo) > 8 * 1024 * 1024:
        return _no("El archivo pesa más de 8 MB", 400)
    r = importar_extracto(caja_id, f.filename or "extracto.csv", crudo)
    if r.get("ok"):
        _auditar("Importó un extracto bancario", r.get("msg"))
    return _resp(r)


@bp_teso.route("/api/teso/conciliacion/<int:caja_id>/auto", methods=["POST"])
def api_auto(caja_id):
    if not _puede(True):
        return _no()
    n = conciliar_auto(caja_id)
    return jsonify({"ok": True, "conciliadas": n, "msg": f"{n} conciliadas" if n else "No encontré más coincidencias seguras"})


@bp_teso.route("/api/teso/extracto/<int:eid>/<accion>", methods=["POST"])
def api_extracto_accion(eid, accion):
    if not _puede(True):
        return _no()
    d = request.get_json(silent=True) or {}
    if accion == "conciliar":
        return _resp(conciliar_manual(eid, int(_num(d.get("mov_id")))))
    if accion == "desconciliar":
        return _resp(desconciliar(eid))
    if accion == "ignorar":
        return _resp(ignorar(eid))
    if accion == "crear":
        return _resp(crear_desde_linea(eid, d))
    return _no("No existe", 404)


@bp_teso.route("/api/teso/cuenta/<int:caja_id>", methods=["PATCH"])
def api_cuenta(caja_id):
    """Datos de banco de una caja: tipo (caja/banco), banco, número, mínimo."""
    if session.get("rol") != "admin":
        return _no("Solo administración")
    d = request.get_json(silent=True) or {}
    campos, vals = [], []
    if "tipo" in d:
        campos.append("tipo=?"); vals.append("banco" if d["tipo"] == "banco" else "caja")
    for k, n in (("banco", 60), ("numero_cuenta", 40)):
        if k in d:
            campos.append(f"{k}=?"); vals.append(_txt(d[k], n))
    if "saldo_minimo" in d:
        campos.append("saldo_minimo=?"); vals.append(max(0, round(_num(d["saldo_minimo"]))))
    if not campos:
        return jsonify({"ok": True})
    conn = get_connection()
    try:
        conn.execute(f"UPDATE cajas SET {', '.join(campos)} WHERE id=?", vals + [caja_id])
        conn.commit()
    finally:
        conn.close()
    return jsonify({"ok": True, "msg": "Cuenta actualizada"})


@bp_teso.route("/api/teso/config", methods=["POST"])
def api_config():
    if session.get("rol") != "admin":
        return _no("Solo administración")
    return jsonify({"ok": True, "config": guardar_config(request.get_json(silent=True) or {}), "msg": "Guardado"})
