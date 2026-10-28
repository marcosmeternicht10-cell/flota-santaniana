"""Módulo contable de La Santaniana.

Cuentas por pagar, comprobantes adjuntos, órdenes de pago con autorización,
caja/bancos, movimientos y exportación. Usa la capa db_compat para funcionar
tanto con SQLite local como con PostgreSQL en la nube.
"""
import base64
import csv
import io
import json
import os
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

from flask import Blueprint, abort, jsonify, make_response, render_template, request, send_file, session

from db_compat import get_connection

bp_contabilidad = Blueprint("contabilidad", __name__)
ROLES_CONTABILIDAD = ("admin", "contabilidad")
MAX_ADJUNTO = 8 * 1024 * 1024
EXTENSIONES_ADMITIDAS = {"pdf", "png", "jpg", "jpeg", "webp", "xlsx", "xls", "docx", "doc", "csv"}
ESTADOS_FACTURA = {"pendiente", "en_pago", "pagada", "anulada"}


def _usuario():
    return session.get("nombre") or session.get("usuario") or "Usuario"


def _rol():
    return session.get("rol") or ""


def _puede_leer():
    if _rol() in ("admin", "contabilidad", "auditor"):
        return True
    if _rol() == "taller":
        try:
            from database import secciones_de_usuario
            return "contabilidad" in (secciones_de_usuario(session.get("usuario_id")) or [])
        except Exception:
            return False
    return False


def _puede_editar():
    if _rol() in ROLES_CONTABILIDAD:
        return True
    if _rol() == "taller":
        try:
            from database import secciones_de_usuario
            secs = secciones_de_usuario(session.get("usuario_id")) or []
            return "contabilidad" in secs and "contabilidad_editar" in secs
        except Exception:
            return False
    return False


def _auditar(accion, detalle="", referencia=""):
    try:
        from database import registrar_auditoria
        registrar_auditoria(usuario=_usuario(), rol=_rol(), accion=accion,
                            categoria="Contabilidad", detalle=detalle, referencia=referencia)
    except Exception:
        pass


def init_contabilidad(app=None):
    """Inicializa las tablas contables sin modificar datos de otros módulos."""
    conn = get_connection()
    try:
        c = conn.cursor()
        c.execute("""CREATE TABLE IF NOT EXISTS conta_facturas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            tipo TEXT NOT NULL DEFAULT 'factura',
            numero TEXT NOT NULL DEFAULT '',
            proveedor TEXT NOT NULL DEFAULT '',
            ruc TEXT NOT NULL DEFAULT '',
            concepto TEXT NOT NULL DEFAULT '',
            fecha_emision TEXT NOT NULL DEFAULT '',
            fecha_vencimiento TEXT NOT NULL DEFAULT '',
            monto REAL NOT NULL DEFAULT 0,
            moneda TEXT NOT NULL DEFAULT 'PYG',
            estado TEXT NOT NULL DEFAULT 'pendiente',
            centro_costo TEXT NOT NULL DEFAULT '',
            observacion TEXT NOT NULL DEFAULT '',
            creado_por TEXT NOT NULL DEFAULT '',
            creado_en TEXT NOT NULL DEFAULT '',
            actualizado_en TEXT NOT NULL DEFAULT ''
        )""")
        c.execute("""CREATE TABLE IF NOT EXISTS conta_adjuntos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            factura_id INTEGER,
            orden_pago_id INTEGER,
            nombre_original TEXT NOT NULL,
            mime_type TEXT NOT NULL DEFAULT 'application/octet-stream',
            tamano INTEGER NOT NULL DEFAULT 0,
            contenido_base64 TEXT NOT NULL,
            subido_por TEXT NOT NULL DEFAULT '',
            creado_en TEXT NOT NULL DEFAULT ''
        )""")
        c.execute("""CREATE TABLE IF NOT EXISTS conta_cuentas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            nombre TEXT NOT NULL,
            tipo TEXT NOT NULL DEFAULT 'banco',
            banco TEXT NOT NULL DEFAULT '',
            moneda TEXT NOT NULL DEFAULT 'PYG',
            saldo_inicial REAL NOT NULL DEFAULT 0,
            activa INTEGER NOT NULL DEFAULT 1,
            creado_por TEXT NOT NULL DEFAULT '',
            creado_en TEXT NOT NULL DEFAULT ''
        )""")
        c.execute("""CREATE TABLE IF NOT EXISTS conta_ordenes_pago (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            numero TEXT NOT NULL DEFAULT '',
            beneficiario TEXT NOT NULL DEFAULT '',
            concepto TEXT NOT NULL DEFAULT '',
            fecha TEXT NOT NULL DEFAULT '',
            monto REAL NOT NULL DEFAULT 0,
            moneda TEXT NOT NULL DEFAULT 'PYG',
            metodo_pago TEXT NOT NULL DEFAULT '',
            cuenta_id INTEGER,
            referencia_pago TEXT NOT NULL DEFAULT '',
            estado TEXT NOT NULL DEFAULT 'pendiente',
            solicitado_por TEXT NOT NULL DEFAULT '',
            autorizado_por TEXT NOT NULL DEFAULT '',
            fecha_autorizacion TEXT NOT NULL DEFAULT '',
            pagado_por TEXT NOT NULL DEFAULT '',
            fecha_pago TEXT NOT NULL DEFAULT '',
            observacion TEXT NOT NULL DEFAULT '',
            creado_en TEXT NOT NULL DEFAULT ''
        )""")
        c.execute("""CREATE TABLE IF NOT EXISTS conta_op_facturas (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            orden_pago_id INTEGER NOT NULL,
            factura_id INTEGER NOT NULL,
            monto_aplicado REAL NOT NULL DEFAULT 0
        )""")
        c.execute("""CREATE TABLE IF NOT EXISTS conta_movimientos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            cuenta_id INTEGER NOT NULL,
            tipo TEXT NOT NULL,
            monto REAL NOT NULL DEFAULT 0,
            fecha TEXT NOT NULL DEFAULT '',
            metodo_pago TEXT NOT NULL DEFAULT '',
            referencia TEXT NOT NULL DEFAULT '',
            orden_pago_id INTEGER,
            descripcion TEXT NOT NULL DEFAULT '',
            creado_por TEXT NOT NULL DEFAULT '',
            creado_en TEXT NOT NULL DEFAULT ''
        )""")
        c.execute("CREATE INDEX IF NOT EXISTS idx_conta_facturas_estado ON conta_facturas(estado)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_conta_op_estado ON conta_ordenes_pago(estado)")
        c.execute("CREATE INDEX IF NOT EXISTS idx_conta_mov_fecha ON conta_movimientos(fecha)")
        conn.commit()
    finally:
        conn.close()


def _dict(row):
    return dict(row) if row is not None else None


def _fetchall(sql, params=()):
    conn = get_connection()
    try:
        return [_dict(r) for r in conn.execute(sql, params).fetchall()]
    finally:
        conn.close()


def _fetchone(sql, params=()):
    conn = get_connection()
    try:
        return _dict(conn.execute(sql, params).fetchone())
    finally:
        conn.close()


def _texto(valor, limite=500):
    return str(valor or "").strip()[:limite]


def _monto(valor):
    try:
        n = Decimal(str(valor).replace(",", "."))
        if not n.is_finite() or n <= 0:
            raise ValueError
        return round(float(n), 2)
    except (InvalidOperation, ValueError, TypeError):
        raise ValueError("El importe debe ser un número mayor que cero.")


def _fecha(valor, opcional=False):
    valor = _texto(valor, 20)
    if not valor and opcional:
        return ""
    try:
        return date.fromisoformat(valor).isoformat()
    except ValueError:
        raise ValueError("Ingresá una fecha válida.")


def _saldo_factura(factura_id, conn=None):
    own = conn is None
    conn = conn or get_connection()
    try:
        row = conn.execute("""SELECT f.monto,
            COALESCE((SELECT SUM(ofa.monto_aplicado)
                      FROM conta_op_facturas ofa
                      JOIN conta_ordenes_pago op ON op.id=ofa.orden_pago_id
                      WHERE ofa.factura_id=f.id AND op.estado='pagada'), 0) AS pagado
            FROM conta_facturas f WHERE f.id=?""", (factura_id,)).fetchone()
        if not row:
            return 0
        return round(max(0, float(row["monto"] or 0) - float(row["pagado"] or 0)), 2)
    finally:
        if own:
            conn.close()


def _numero_op(conn, op_id):
    numero = f"OP-{op_id:06d}"
    conn.execute("UPDATE conta_ordenes_pago SET numero=? WHERE id=?", (numero, op_id))
    return numero


@bp_contabilidad.route("/contabilidad")
def pagina_contabilidad():
    if not _puede_leer():
        abort(403)
    return render_template("contabilidad.html",
                           usuario=_usuario(), rol=_rol(),
                           puede_editar=_puede_editar())


@bp_contabilidad.route("/api/contabilidad/resumen")
def api_conta_resumen():
    if not _puede_leer():
        return jsonify({"error": "Sin permiso"}), 403
    pendientes = _fetchone("""SELECT COUNT(*) AS n, COALESCE(SUM(monto),0) AS total
                              FROM conta_facturas WHERE estado IN ('pendiente','en_pago')""") or {}
    vencidas = _fetchone("""SELECT COUNT(*) AS n, COALESCE(SUM(monto),0) AS total
                            FROM conta_facturas WHERE estado IN ('pendiente','en_pago')
                            AND fecha_vencimiento<>'' AND fecha_vencimiento<?""", (date.today().isoformat(),)) or {}
    por_autorizar = _fetchone("""SELECT COUNT(*) AS n, COALESCE(SUM(monto),0) AS total
                                 FROM conta_ordenes_pago WHERE estado='pendiente'""") or {}
    pagado = _fetchone("""SELECT COUNT(*) AS n, COALESCE(SUM(monto),0) AS total
                          FROM conta_ordenes_pago WHERE estado='pagada'
                          AND substr(fecha_pago,1,7)=?""", (date.today().strftime("%Y-%m"),)) or {}
    cuentas = _fetchall("""SELECT c.id,c.nombre,c.tipo,c.banco,c.moneda,c.saldo_inicial,c.activa,
        c.saldo_inicial + COALESCE(SUM(CASE WHEN m.tipo='ingreso' THEN m.monto
                                            WHEN m.tipo='egreso' THEN -m.monto ELSE 0 END),0) AS saldo
        FROM conta_cuentas c LEFT JOIN conta_movimientos m ON m.cuenta_id=c.id
        WHERE c.activa=1 GROUP BY c.id,c.nombre,c.tipo,c.banco,c.moneda,c.saldo_inicial,c.activa
        ORDER BY c.nombre""")
    return jsonify({
        "facturas_pendientes": int(pendientes.get("n") or 0),
        "monto_pendiente": float(pendientes.get("total") or 0),
        "facturas_vencidas": int(vencidas.get("n") or 0),
        "monto_vencido": float(vencidas.get("total") or 0),
        "op_pendientes": int(por_autorizar.get("n") or 0),
        "op_pendientes_monto": float(por_autorizar.get("total") or 0),
        "pagado_mes": float(pagado.get("total") or 0),
        "ordenes_pagadas_mes": int(pagado.get("n") or 0),
        "cuentas": cuentas
    })


@bp_contabilidad.route("/api/contabilidad/facturas", methods=["GET", "POST"])
def api_conta_facturas():
    if request.method == "GET":
        if not _puede_leer():
            return jsonify({"error": "Sin permiso"}), 403
        estado = _texto(request.args.get("estado"), 30)
        buscar = _texto(request.args.get("buscar"), 120)
        sql = """SELECT f.*,
            COALESCE((SELECT SUM(ofa.monto_aplicado) FROM conta_op_facturas ofa
                      JOIN conta_ordenes_pago op ON op.id=ofa.orden_pago_id
                      WHERE ofa.factura_id=f.id AND op.estado='pagada'),0) AS monto_pagado
            FROM conta_facturas f WHERE 1=1"""
        params = []
        if estado:
            sql += " AND f.estado=?"
            params.append(estado)
        if buscar:
            sql += " AND (LOWER(f.proveedor) LIKE ? OR LOWER(f.numero) LIKE ? OR LOWER(f.concepto) LIKE ? OR LOWER(f.ruc) LIKE ?)"
            q = "%" + buscar.lower() + "%"
            params.extend([q, q, q, q])
        sql += " ORDER BY CASE WHEN f.fecha_vencimiento='' THEN 1 ELSE 0 END, f.fecha_vencimiento, f.id DESC LIMIT 500"
        rows = _fetchall(sql, tuple(params))
        for row in rows:
            row["saldo"] = round(max(0, float(row["monto"] or 0) - float(row["monto_pagado"] or 0)), 2)
        return jsonify(rows)
    if not _puede_editar():
        return jsonify({"error": "No tenés permiso para registrar comprobantes"}), 403
    d = request.get_json(silent=True) or {}
    try:
        proveedor = _texto(d.get("proveedor"), 180)
        concepto = _texto(d.get("concepto"), 300)
        if not proveedor or not concepto:
            raise ValueError("Proveedor/beneficiario y concepto son obligatorios.")
        monto = _monto(d.get("monto"))
        fecha_emision = _fecha(d.get("fecha_emision"), opcional=True)
        fecha_vencimiento = _fecha(d.get("fecha_vencimiento"), opcional=True)
        tipo = _texto(d.get("tipo") or "factura", 30)
        if tipo not in ("factura", "gasto", "recibo", "otro"):
            tipo = "otro"
        moneda = _texto(d.get("moneda") or "PYG", 8).upper()
        if moneda not in ("PYG", "USD", "ARS", "BRL"):
            moneda = "PYG"
        conn = get_connection()
        try:
            now = datetime.now().isoformat(timespec="seconds")
            cur = conn.execute("""INSERT INTO conta_facturas
                (tipo,numero,proveedor,ruc,concepto,fecha_emision,fecha_vencimiento,monto,moneda,estado,centro_costo,observacion,creado_por,creado_en,actualizado_en)
                VALUES (?,?,?,?,?,?,?,?,?,'pendiente',?,?,?,?,?)""",
                (tipo, _texto(d.get("numero"), 100), proveedor, _texto(d.get("ruc"), 40), concepto,
                 fecha_emision, fecha_vencimiento, monto, moneda, _texto(d.get("centro_costo"), 120),
                 _texto(d.get("observacion"), 1500), _usuario(), now, now))
            fid = cur.lastrowid
            conn.commit()
        finally:
            conn.close()
        _auditar("Registró un comprobante", f"{proveedor} · {_texto(d.get('numero'),100)} · {monto} {moneda}", f"conta:factura:{fid}")
        return jsonify({"ok": True, "id": fid}), 201
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    except Exception:
        return jsonify({"error": "No se pudo guardar el comprobante. Revisá los datos e intentá de nuevo."}), 500


@bp_contabilidad.route("/api/contabilidad/facturas/<int:factura_id>/adjuntos", methods=["GET", "POST"])
def api_conta_adjuntos(factura_id):
    if request.method == "GET":
        if not _puede_leer():
            return jsonify({"error": "Sin permiso"}), 403
        if not _fetchone("SELECT id FROM conta_facturas WHERE id=?", (factura_id,)):
            return jsonify({"error": "Comprobante no encontrado"}), 404
        return jsonify(_fetchall("SELECT id,nombre_original,mime_type,tamano,subido_por,creado_en FROM conta_adjuntos WHERE factura_id=? ORDER BY id", (factura_id,)))
    if not _puede_editar():
        return jsonify({"error": "Sin permiso"}), 403
    if not _fetchone("SELECT id FROM conta_facturas WHERE id=?", (factura_id,)):
        return jsonify({"error": "Comprobante no encontrado"}), 404
    archivos = request.files.getlist("archivos") or request.files.getlist("archivo")
    if not archivos or not any(a.filename for a in archivos):
        return jsonify({"error": "Elegí al menos un archivo."}), 400
    guardados = 0
    conn = get_connection()
    try:
        for archivo in archivos:
            if not archivo or not archivo.filename:
                continue
            nombre = os.path.basename(archivo.filename.replace("\\", "/"))
            ext = nombre.rsplit(".", 1)[-1].lower() if "." in nombre else ""
            if ext not in EXTENSIONES_ADMITIDAS:
                return jsonify({"error": f"Tipo de archivo no admitido: {nombre}"}), 400
            contenido = archivo.read(MAX_ADJUNTO + 1)
            if len(contenido) > MAX_ADJUNTO:
                return jsonify({"error": f"{nombre} supera el límite de 8 MB."}), 400
            if not contenido:
                continue
            conn.execute("""INSERT INTO conta_adjuntos
                (factura_id,orden_pago_id,nombre_original,mime_type,tamano,contenido_base64,subido_por,creado_en)
                VALUES (?,NULL,?,?,?,?,?,?)""",
                (factura_id, nombre[:240], archivo.mimetype or "application/octet-stream",
                 len(contenido), base64.b64encode(contenido).decode("ascii"), _usuario(),
                 datetime.now().isoformat(timespec="seconds")))
            guardados += 1
        conn.commit()
    finally:
        conn.close()
    _auditar("Adjuntó documentación a un comprobante", f"{guardados} archivo(s)", f"conta:factura:{factura_id}")
    return jsonify({"ok": True, "guardados": guardados})


@bp_contabilidad.route("/api/contabilidad/adjuntos/<int:adjunto_id>")
def api_conta_descargar_adjunto(adjunto_id):
    if not _puede_leer():
        return jsonify({"error": "Sin permiso"}), 403
    row = _fetchone("SELECT nombre_original,mime_type,contenido_base64 FROM conta_adjuntos WHERE id=?", (adjunto_id,))
    if not row:
        return jsonify({"error": "Adjunto no encontrado"}), 404
    try:
        data = base64.b64decode(row["contenido_base64"])
    except Exception:
        return jsonify({"error": "El archivo guardado no se puede leer"}), 500
    return send_file(io.BytesIO(data), mimetype=row["mime_type"] or "application/octet-stream",
                     as_attachment=True, download_name=row["nombre_original"])


@bp_contabilidad.route("/api/contabilidad/ordenes", methods=["GET", "POST"])
def api_conta_ordenes():
    if request.method == "GET":
        if not _puede_leer():
            return jsonify({"error": "Sin permiso"}), 403
        estado = _texto(request.args.get("estado"), 30)
        sql = """SELECT op.*, c.nombre AS cuenta_nombre,
                 (SELECT COUNT(*) FROM conta_op_facturas ofa WHERE ofa.orden_pago_id=op.id) AS comprobantes
                 FROM conta_ordenes_pago op LEFT JOIN conta_cuentas c ON c.id=op.cuenta_id WHERE 1=1"""
        params = []
        if estado:
            sql += " AND op.estado=?"
            params.append(estado)
        sql += " ORDER BY op.id DESC LIMIT 500"
        rows = _fetchall(sql, tuple(params))
        for row in rows:
            row["facturas"] = _fetchall("""SELECT f.id,f.numero,f.proveedor,f.concepto,ofa.monto_aplicado
                FROM conta_op_facturas ofa JOIN conta_facturas f ON f.id=ofa.factura_id
                WHERE ofa.orden_pago_id=? ORDER BY f.id""", (row["id"],))
        return jsonify(rows)
    if not _puede_editar():
        return jsonify({"error": "Sin permiso para crear órdenes de pago"}), 403
    d = request.get_json(silent=True) or {}
    ids = d.get("factura_ids") or []
    try:
        ids = list(dict.fromkeys(int(x) for x in ids))
    except (TypeError, ValueError):
        return jsonify({"error": "Seleccioná comprobantes válidos."}), 400
    if not ids:
        return jsonify({"error": "Seleccioná al menos un comprobante pendiente."}), 400
    beneficiario = _texto(d.get("beneficiario"), 180)
    concepto = _texto(d.get("concepto"), 300)
    if not beneficiario or not concepto:
        return jsonify({"error": "Beneficiario y concepto son obligatorios."}), 400
    conn = get_connection()
    try:
        placeholders = ",".join("?" for _ in ids)
        facturas = conn.execute(f"SELECT * FROM conta_facturas WHERE id IN ({placeholders})", tuple(ids)).fetchall()
        if len(facturas) != len(ids):
            return jsonify({"error": "Uno o más comprobantes no existen."}), 404
        moneda = facturas[0]["moneda"]
        if any(f["moneda"] != moneda for f in facturas):
            return jsonify({"error": "No se pueden mezclar monedas en una misma orden de pago."}), 400
        monto_total = 0
        for f in facturas:
            if f["estado"] not in ("pendiente",):
                return jsonify({"error": f"El comprobante #{f['id']} no está disponible para pagar."}), 400
            saldo = _saldo_factura(f["id"], conn)
            if saldo <= 0:
                return jsonify({"error": f"El comprobante #{f['id']} ya está saldado."}), 400
            monto_total += saldo
        now = datetime.now().isoformat(timespec="seconds")
        fecha = _fecha(d.get("fecha") or date.today().isoformat())
        cur = conn.execute("""INSERT INTO conta_ordenes_pago
            (numero,beneficiario,concepto,fecha,monto,moneda,metodo_pago,cuenta_id,referencia_pago,estado,solicitado_por,autorizado_por,fecha_autorizacion,pagado_por,fecha_pago,observacion,creado_en)
            VALUES ('',?,?,?,?,?,'',NULL,'','pendiente',?,'','','','',?,?)""",
            (beneficiario, concepto, fecha, round(monto_total, 2), moneda, _usuario(),
             _texto(d.get("observacion"), 1500), now))
        op_id = cur.lastrowid
        numero = _numero_op(conn, op_id)
        for f in facturas:
            saldo = _saldo_factura(f["id"], conn)
            conn.execute("INSERT INTO conta_op_facturas (orden_pago_id,factura_id,monto_aplicado) VALUES (?,?,?)",
                         (op_id, f["id"], saldo))
            conn.execute("UPDATE conta_facturas SET estado='en_pago',actualizado_en=? WHERE id=?",
                         (now, f["id"]))
        conn.commit()
    except ValueError as e:
        conn.rollback()
        return jsonify({"error": str(e)}), 400
    except Exception:
        conn.rollback()
        return jsonify({"error": "No se pudo crear la orden de pago."}), 500
    finally:
        conn.close()
    _auditar("Creó orden de pago", f"{numero} · {beneficiario} · {monto_total} {moneda}", f"conta:op:{op_id}")
    return jsonify({"ok": True, "id": op_id, "numero": numero, "monto": round(monto_total, 2)}), 201


@bp_contabilidad.route("/api/contabilidad/ordenes/<int:op_id>/autorizar", methods=["POST"])
def api_conta_autorizar(op_id):
    if not _puede_editar():
        return jsonify({"error": "Sin permiso para autorizar órdenes"}), 403
    conn = get_connection()
    try:
        op = conn.execute("SELECT * FROM conta_ordenes_pago WHERE id=?", (op_id,)).fetchone()
        if not op:
            return jsonify({"error": "Orden no encontrada"}), 404
        if op["estado"] != "pendiente":
            return jsonify({"error": "Solo se pueden autorizar órdenes pendientes."}), 400
        if op["solicitado_por"] == _usuario() and _rol() != "admin":
            return jsonify({"error": "Por control interno, otra persona debe autorizar una orden que vos creaste."}), 403
        now = datetime.now().isoformat(timespec="seconds")
        conn.execute("UPDATE conta_ordenes_pago SET estado='autorizada',autorizado_por=?,fecha_autorizacion=? WHERE id=?",
                     (_usuario(), now, op_id))
        conn.commit()
    finally:
        conn.close()
    _auditar("Autorizó orden de pago", f"OP-{op_id:06d}", f"conta:op:{op_id}")
    return jsonify({"ok": True})


@bp_contabilidad.route("/api/contabilidad/ordenes/<int:op_id>/rechazar", methods=["POST"])
def api_conta_rechazar(op_id):
    if not _puede_editar():
        return jsonify({"error": "Sin permiso para rechazar órdenes"}), 403
    conn = get_connection()
    try:
        op = conn.execute("SELECT * FROM conta_ordenes_pago WHERE id=?", (op_id,)).fetchone()
        if not op:
            return jsonify({"error": "Orden no encontrada"}), 404
        if op["estado"] != "pendiente":
            return jsonify({"error": "Solo se pueden rechazar órdenes pendientes."}), 400
        d = request.get_json(silent=True) or {}
        motivo = _texto(d.get("motivo"), 1000)
        if not motivo:
            return jsonify({"error": "Indicá el motivo del rechazo."}), 400
        conn.execute("UPDATE conta_ordenes_pago SET estado='rechazada',observacion=? WHERE id=?", (motivo, op_id))
        conn.execute("""UPDATE conta_facturas SET estado='pendiente',actualizado_en=?
                        WHERE id IN (SELECT factura_id FROM conta_op_facturas WHERE orden_pago_id=?)
                        AND estado='en_pago'""", (datetime.now().isoformat(timespec="seconds"), op_id))
        conn.commit()
    finally:
        conn.close()
    _auditar("Rechazó orden de pago", f"OP-{op_id:06d}: {motivo}", f"conta:op:{op_id}")
    return jsonify({"ok": True})


@bp_contabilidad.route("/api/contabilidad/ordenes/<int:op_id>/pagar", methods=["POST"])
def api_conta_pagar(op_id):
    if not _puede_editar():
        return jsonify({"error": "Sin permiso para registrar pagos"}), 403
    d = request.get_json(silent=True) or {}
    try:
        cuenta_id = int(d.get("cuenta_id") or 0)
        fecha_pago = _fecha(d.get("fecha_pago") or date.today().isoformat())
        metodo = _texto(d.get("metodo_pago"), 40)
        referencia = _texto(d.get("referencia_pago"), 120)
        if not cuenta_id or not metodo:
            raise ValueError("Elegí la cuenta de origen y el método de pago.")
    except ValueError as e:
        return jsonify({"error": str(e)}), 400
    conn = get_connection()
    try:
        op = conn.execute("SELECT * FROM conta_ordenes_pago WHERE id=?", (op_id,)).fetchone()
        if not op:
            return jsonify({"error": "Orden no encontrada"}), 404
        if op["estado"] != "autorizada":
            return jsonify({"error": "La orden debe estar autorizada antes de registrar el pago."}), 400
        cuenta = conn.execute("SELECT * FROM conta_cuentas WHERE id=? AND activa=1", (cuenta_id,)).fetchone()
        if not cuenta:
            return jsonify({"error": "La cuenta elegida no existe o está inactiva."}), 400
        if cuenta["moneda"] != op["moneda"]:
            return jsonify({"error": "La moneda de la cuenta no coincide con la orden de pago."}), 400
        if conn.execute("SELECT id FROM conta_movimientos WHERE orden_pago_id=?", (op_id,)).fetchone():
            return jsonify({"error": "Esta orden ya tiene un movimiento registrado."}), 409
        now = datetime.now().isoformat(timespec="seconds")
        conn.execute("""INSERT INTO conta_movimientos
            (cuenta_id,tipo,monto,fecha,metodo_pago,referencia,orden_pago_id,descripcion,creado_por,creado_en)
            VALUES (?,'egreso',?,?,?,?,?,?,?,?)""",
            (cuenta_id, float(op["monto"]), fecha_pago, metodo, referencia, op_id,
             f"{op['numero']} · {op['beneficiario']} · {op['concepto']}", _usuario(), now))
        conn.execute("""UPDATE conta_ordenes_pago
            SET estado='pagada',cuenta_id=?,metodo_pago=?,referencia_pago=?,pagado_por=?,fecha_pago=?
            WHERE id=?""", (cuenta_id, metodo, referencia, _usuario(), fecha_pago, op_id))
        conn.execute("""UPDATE conta_facturas SET estado='pagada',actualizado_en=?
            WHERE id IN (SELECT factura_id FROM conta_op_facturas WHERE orden_pago_id=?)""", (now, op_id))
        conn.commit()
    except Exception as e:
        conn.rollback()
        if isinstance(e, ValueError):
            return jsonify({"error": str(e)}), 400
        return jsonify({"error": "No se pudo registrar el pago."}), 500
    finally:
        conn.close()
    _auditar("Registró pago de orden", f"OP-{op_id:06d} · {op['monto']} {op['moneda']}", f"conta:op:{op_id}")
    return jsonify({"ok": True})


@bp_contabilidad.route("/api/contabilidad/cuentas", methods=["GET", "POST"])
def api_conta_cuentas():
    if request.method == "GET":
        if not _puede_leer():
            return jsonify({"error": "Sin permiso"}), 403
        return jsonify(_fetchall("""SELECT c.id,c.nombre,c.tipo,c.banco,c.moneda,c.saldo_inicial,c.activa,
            c.saldo_inicial + COALESCE(SUM(CASE WHEN m.tipo='ingreso' THEN m.monto
                WHEN m.tipo='egreso' THEN -m.monto ELSE 0 END),0) AS saldo
            FROM conta_cuentas c LEFT JOIN conta_movimientos m ON m.cuenta_id=c.id
            GROUP BY c.id,c.nombre,c.tipo,c.banco,c.moneda,c.saldo_inicial,c.activa ORDER BY c.nombre"""))
    if not _puede_editar():
        return jsonify({"error": "Sin permiso para administrar cuentas"}), 403
    d = request.get_json(silent=True) or {}
    nombre = _texto(d.get("nombre"), 120)
    tipo = _texto(d.get("tipo") or "banco", 20)
    moneda = _texto(d.get("moneda") or "PYG", 8).upper()
    if not nombre or tipo not in ("caja", "banco") or moneda not in ("PYG", "USD", "ARS", "BRL"):
        return jsonify({"error": "Completá el nombre, tipo de cuenta y moneda válidos."}), 400
    try:
        saldo = float(Decimal(str(d.get("saldo_inicial") or 0).replace(",", ".")))
        if not Decimal(str(saldo)).is_finite():
            raise ValueError
    except Exception:
        return jsonify({"error": "El saldo inicial no es válido."}), 400
    conn = get_connection()
    try:
        cur = conn.execute("""INSERT INTO conta_cuentas
            (nombre,tipo,banco,moneda,saldo_inicial,activa,creado_por,creado_en)
            VALUES (?,?,?,?,?,1,?,?)""",
            (nombre, tipo, _texto(d.get("banco"), 120), moneda, saldo, _usuario(), datetime.now().isoformat(timespec="seconds")))
        cid = cur.lastrowid
        conn.commit()
    finally:
        conn.close()
    _auditar("Creó cuenta contable", f"{nombre} · {tipo} · {moneda}", f"conta:cuenta:{cid}")
    return jsonify({"ok": True, "id": cid}), 201


@bp_contabilidad.route("/api/contabilidad/movimientos")
def api_conta_movimientos():
    if not _puede_leer():
        return jsonify({"error": "Sin permiso"}), 403
    desde = _texto(request.args.get("desde"), 10)
    hasta = _texto(request.args.get("hasta"), 10)
    sql = """SELECT m.*,c.nombre AS cuenta_nombre,op.numero AS orden_numero
             FROM conta_movimientos m JOIN conta_cuentas c ON c.id=m.cuenta_id
             LEFT JOIN conta_ordenes_pago op ON op.id=m.orden_pago_id WHERE 1=1"""
    params = []
    if desde:
        sql += " AND m.fecha>=?"
        params.append(desde)
    if hasta:
        sql += " AND m.fecha<=?"
        params.append(hasta)
    sql += " ORDER BY m.fecha DESC,m.id DESC LIMIT 1000"
    return jsonify(_fetchall(sql, tuple(params)))


@bp_contabilidad.route("/api/contabilidad/exportar.xlsx")
def api_conta_exportar():
    if not _puede_leer():
        return jsonify({"error": "Sin permiso"}), 403
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        return jsonify({"error": "Falta instalar openpyxl. Actualizá las dependencias del sistema."}), 500
    facturas = _fetchall("SELECT id,tipo,numero,proveedor,ruc,concepto,fecha_emision,fecha_vencimiento,monto,moneda,estado,centro_costo,observacion,creado_por,creado_en FROM conta_facturas ORDER BY id DESC")
    ordenes = _fetchall("SELECT numero,beneficiario,concepto,fecha,monto,moneda,metodo_pago,estado,solicitado_por,autorizado_por,fecha_autorizacion,pagado_por,fecha_pago,referencia_pago,observacion FROM conta_ordenes_pago ORDER BY id DESC")
    movimientos = _fetchall("""SELECT m.fecha,m.tipo,m.descripcion,m.monto,c.nombre AS cuenta,m.metodo_pago,m.referencia,op.numero AS orden_pago,m.creado_por
        FROM conta_movimientos m JOIN conta_cuentas c ON c.id=m.cuenta_id LEFT JOIN conta_ordenes_pago op ON op.id=m.orden_pago_id ORDER BY m.fecha DESC,m.id DESC""")
    cuentas = _fetchall("""SELECT c.nombre,c.tipo,c.banco,c.moneda,c.saldo_inicial,
        c.saldo_inicial+COALESCE(SUM(CASE WHEN m.tipo='ingreso' THEN m.monto WHEN m.tipo='egreso' THEN -m.monto ELSE 0 END),0) AS saldo
        FROM conta_cuentas c LEFT JOIN conta_movimientos m ON m.cuenta_id=c.id
        GROUP BY c.id,c.nombre,c.tipo,c.banco,c.moneda,c.saldo_inicial ORDER BY c.nombre""")
    wb = Workbook()
    for title, rows in (("Comprobantes", facturas), ("Ordenes de pago", ordenes), ("Movimientos", movimientos), ("Caja y bancos", cuentas)):
        ws = wb.active if title == "Comprobantes" else wb.create_sheet()
        ws.title = title
        if not rows:
            ws.append(["Sin registros"])
            continue
        headers = list(rows[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="17324D")
            cell.alignment = Alignment(vertical="center")
        for row in rows:
            ws.append([row.get(h) for h in headers])
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        for col in ws.columns:
            letter = col[0].column_letter
            width = min(38, max(12, max(len(str(cell.value or "")) for cell in col[:100]) + 2))
            ws.column_dimensions[letter].width = width
    out = io.BytesIO()
    wb.save(out)
    out.seek(0)
    return send_file(out, as_attachment=True, download_name=f"contabilidad_{date.today().isoformat()}.xlsx",
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
