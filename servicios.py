"""
servicios.py — Servicios: lo que salió de verdad, sacado del tráfico (La Santaniana)

Cada salida del tráfico de un día que ya llegó es un servicio. No se carga nada
aparte: lo que la encargada arma en Tráfico es la fuente. Con eso se arman los
indicadores para entender dónde estamos parados y qué mejorar:

  - Cuántos servicios salen, por día, por corredor y por horario.
  - Refuerzos: cuántos y dónde se repiten (la demanda que pide un horario fijo).
  - La flota: cuántos coches trabajan, cuáles cargan más y cuáles están parados.
  - La tripulación: cuántos servicios hace cada uno (para repartir mejor).
  - Cómo se arma el tráfico: con cuánta anticipación se publica, cuántas
    versiones hay por día y cuántas salidas se mandan A CONFIRMAR.
  - Km y asientos ofrecidos: con la distancia de cada tramo (se carga una vez)
    y los asientos de cada coche (de la ficha del vehículo).

Además arma un Excel con el detalle para liquidar o revisar.

Enganche en app.py:
    from servicios import bp_servicios, init_servicios
    init_servicios()
    app.register_blueprint(bp_servicios)
"""

import io
import json
import statistics
from datetime import date, datetime, timedelta

from flask import Blueprint, request, jsonify, session, send_file
from db_compat import get_connection, USE_POSTGRES
from trafico import tramo, partes_destino, limpiar, _sin_tildes, _ahora, DIAS_SEMANA

bp_servicios = Blueprint("servicios_trafico", __name__)

PK = "SERIAL PRIMARY KEY" if USE_POSTGRES else "INTEGER PRIMARY KEY AUTOINCREMENT"
DIAS_CORTOS = ["LUN", "MAR", "MIÉ", "JUE", "VIE", "SÁB", "DOM"]
MAX_DIAS = 400                       # el período más largo que se puede pedir


def init_servicios():
    conn = get_connection()
    conn.execute(f"""
        CREATE TABLE IF NOT EXISTS servicios_distancias (
            id {PK},
            clave TEXT NOT NULL UNIQUE,      -- ASU|COP o ASU|COP|X 25 (el par ordenado + por dónde va)
            km REAL,
            actualizado TEXT DEFAULT '',
            actualizado_por TEXT DEFAULT ''
        )""")
    conn.commit()
    conn.close()


def _hoy():
    return _ahora()[:10]


# ════════════════════════════════════════════════════════════════════════════
# LOS SERVICIOS DE UN PERÍODO
# ════════════════════════════════════════════════════════════════════════════

def clave_tramo(origen, llega, via=""):
    """El mismo tramo en los dos sentidos es la misma distancia."""
    if not origen or not llega or origen == llega:
        return None
    par = "|".join(sorted([origen, llega]))
    return f"{par}|{via}" if via else par


def nombre_tramo(clave):
    partes = clave.split("|")
    txt = f"{partes[0]} ↔ {partes[1]}"
    return txt + (f" · {partes[2]}" if len(partes) > 2 and partes[2] else "")


def _clave_persona(n):
    return _sin_tildes(limpiar(n))


def _filas(conn, desde, hasta):
    return conn.execute("""
        SELECT s.id, s.lado, s.hora, s.destino, s.coche, s.trip1, s.trip2,
               b.nombre AS corredor, d.fecha, d.id AS dia_id
        FROM trafico_salidas s
        JOIN trafico_bloques b ON b.id = s.bloque_id
        JOIN trafico_dias d ON d.id = s.dia_id
        WHERE d.fecha >= ? AND d.fecha <= ?
        ORDER BY d.fecha, s.hora, s.id""", (desde, hasta)).fetchall()


def leer_servicios(conn, desde, hasta, flota=None, distancias=None):
    """Las salidas del período convertidas en servicios, con lo que hace
    falta para medir: tramo en siglas, si es refuerzo, coche, tripulantes,
    km (si se conoce la distancia) y asientos (si se conoce el coche)."""
    flota = flota if flota is not None else _flota(conn)
    distancias = distancias if distancias is not None else _distancias(conn)
    salida = []
    for f in _filas(conn, desde, hasta):
        corredor = limpiar(f["corredor"])
        origen, llega = tramo(f["destino"], f["lado"], corredor)
        base, extras = partes_destino(f["destino"])
        via = next((e for e in extras if e.startswith("X ")), "")
        clave = clave_tramo(origen, llega, via)
        km = distancias.get(clave) if clave else None
        if km is None and clave and via:                  # sin la distancia "por dónde va", la del par
            km = distancias.get(clave_tramo(origen, llega))
        coche = limpiar(f["coche"])
        coche = coche if coche.isdigit() else ""
        v = flota.get(coche) if coche else None
        trips = [limpiar(t) for t in (f["trip1"], f["trip2"]) if limpiar(t)]
        hora = f["hora"] or ""
        salida.append({
            "id": f["id"], "fecha": f["fecha"], "hora": hora, "lado": f["lado"] if f["lado"] in ("ida", "regreso") else "ida",
            "corredor": corredor, "destino": limpiar(f["destino"]), "tramo_base": base, "extras": extras,
            "origen": origen, "llega": llega, "clave_tramo": clave, "km": km,
            "refuerzo": "REFUERZO" in extras, "coche": coche,
            "asientos": int(v["asientos"] or 0) if v and v.get("asientos") else 0,
            "modelo": f"{v.get('marca') or ''} {v.get('modelo') or ''}".strip() if v else "",
            "trips": trips,
        })
    return salida


def _flota(conn):
    out = {}
    for r in conn.execute("""SELECT id, n_interno, patente, marca, modelo, asientos, tipo,
                                    COALESCE(activo, 1) AS activo FROM vehiculos""").fetchall():
        n = str(r["n_interno"] or "").strip()
        if n:
            out[n] = dict(r)
    return out


def _distancias(conn):
    out = {}
    try:
        for r in conn.execute("SELECT clave, km FROM servicios_distancias").fetchall():
            if r["km"]:
                out[r["clave"]] = float(r["km"])
    except Exception:
        pass
    return out


def _deltas(tot, ant):
    """Contra el período anterior. Lo que se suma (servicios, refuerzos, km,
    asientos) se compara por día con tráfico: si el tráfico empezó a cargarse
    a mitad del período anterior, comparar totales mentiría. Lo que se cuenta
    distinto (coches, tripulantes) solo si los dos períodos tienen datos parejos."""
    da, dp = tot.get("dias") or 0, ant.get("dias") or 0
    out = {}
    for k in ("servicios", "refuerzos", "asientos", "km"):
        a = tot.get(k) / da if da else None
        p = ant.get(k) / dp if dp else None
        out[k] = _delta(a, p) if p else None
    parejo = dp >= 0.8 * da and dp > 0
    for k in ("coches", "tripulantes"):
        out[k] = _delta(tot.get(k), ant.get(k)) if parejo else None
    for k in ("por_dia", "uso_flota", "coches_por_dia"):
        out[k] = _delta(tot.get(k), ant.get(k))
    out["comparable"] = parejo
    return out


def _rango(desde, hasta):
    d0, d1 = date.fromisoformat(desde), date.fromisoformat(hasta)
    return [(d0 + timedelta(days=i)).isoformat() for i in range((d1 - d0).days + 1)]


def _redondo(x, n=1):
    return round(x, n) if x is not None else None


def _delta(ahora, antes):
    """Variación porcentual contra el período anterior (None si no hay base)."""
    if antes in (None, 0) or ahora is None:
        return None
    return round((ahora - antes) / antes * 100, 1)


# ════════════════════════════════════════════════════════════════════════════
# INDICADORES
# ════════════════════════════════════════════════════════════════════════════

def _totales(servs, dias_con_trafico):
    coches = {s["coche"] for s in servs if s["coche"]}
    personas = {_clave_persona(t) for s in servs for t in s["trips"]}
    km_conocidos = [s["km"] for s in servs if s["km"]]
    n = len(servs)
    return {
        "servicios": n,
        "por_dia": round(n / dias_con_trafico, 1) if dias_con_trafico else 0,
        "refuerzos": sum(1 for s in servs if s["refuerzo"]),
        "coches": len(coches),
        "tripulantes": len(personas),
        "asientos": sum(s["asientos"] for s in servs),
        "km": round(sum(km_conocidos)) if km_conocidos else 0,
        "km_cubiertos": len(km_conocidos),          # cuántos servicios tienen la distancia
        "sin_coche": sum(1 for s in servs if not s["coche"]),
        "sin_tripulantes": sum(1 for s in servs if not s["trips"]),
        "dias": dias_con_trafico,
    }


def _serie_diaria(servs, fechas, flota_base_n):
    por = {f: {"fecha": f, "servicios": 0, "refuerzos": 0, "ida": 0, "regreso": 0, "coches": set(),
               "km": 0, "asientos": 0, "con_trafico": False} for f in fechas}
    for s in servs:
        d = por.get(s["fecha"])
        if not d:
            continue
        d["con_trafico"] = True
        d["servicios"] += 1
        d["refuerzos"] += 1 if s["refuerzo"] else 0
        d[s["lado"]] += 1
        if s["coche"]:
            d["coches"].add(s["coche"])
        d["km"] += s["km"] or 0
        d["asientos"] += s["asientos"]
    out = []
    for f in fechas:
        d = por[f]
        n_coches = len(d["coches"])
        out.append({**d, "coches": n_coches, "km": round(d["km"]),
                    "uso": round(n_coches / flota_base_n * 100, 1) if flota_base_n and d["con_trafico"] else None,
                    "dia_semana": date.fromisoformat(f).weekday()})
    return out


def _por_corredor(servs):
    out = {}
    for s in servs:
        c = out.setdefault(s["corredor"], {"corredor": s["corredor"], "servicios": 0, "refuerzos": 0, "ida": 0,
                                           "regreso": 0, "coches": set(), "asientos": 0, "km": 0, "km_cubiertos": 0})
        c["servicios"] += 1
        c["refuerzos"] += 1 if s["refuerzo"] else 0
        c[s["lado"]] += 1
        if s["coche"]:
            c["coches"].add(s["coche"])
        c["asientos"] += s["asientos"]
        if s["km"]:
            c["km"] += s["km"]
            c["km_cubiertos"] += 1
    total = len(servs) or 1
    lista = []
    for c in out.values():
        lista.append({**c, "coches": len(c["coches"]), "km": round(c["km"]),
                      "parte": round(c["servicios"] / total * 100, 1)})
    lista.sort(key=lambda c: (-c["servicios"], c["corredor"]))
    return lista


def _mapa_horario(servs, fechas_con_trafico):
    """Servicios por día de la semana y hora, y cuántos de cada día hubo (para el promedio)."""
    celdas = [[0] * 24 for _ in range(7)]
    for s in servs:
        if not s["hora"]:
            continue
        try:
            h = int(s["hora"][:2])
        except ValueError:
            continue
        celdas[date.fromisoformat(s["fecha"]).weekday()][h] += 1
    dias_semana = [0] * 7
    for f in fechas_con_trafico:
        dias_semana[date.fromisoformat(f).weekday()] += 1
    return {"celdas": celdas, "dias": dias_semana}


def _por_coche(servs, flota):
    out = {}
    for s in servs:
        if not s["coche"]:
            continue
        c = out.setdefault(s["coche"], {"coche": s["coche"], "servicios": 0, "refuerzos": 0, "km": 0, "dias": set(),
                                        "corredores": {}, "ultimo": ""})
        c["servicios"] += 1
        c["refuerzos"] += 1 if s["refuerzo"] else 0
        c["km"] += s["km"] or 0
        c["dias"].add(s["fecha"])
        c["corredores"][s["corredor"]] = c["corredores"].get(s["corredor"], 0) + 1
        c["ultimo"] = max(c["ultimo"], f"{s['fecha']} {s['hora']}")
    lista = []
    for c in out.values():
        v = flota.get(c["coche"]) or {}
        principal = max(c["corredores"].items(), key=lambda x: x[1])[0] if c["corredores"] else ""
        lista.append({"coche": c["coche"], "servicios": c["servicios"], "refuerzos": c["refuerzos"],
                      "km": round(c["km"]), "dias": len(c["dias"]), "corredor": principal,
                      "ultimo": c["ultimo"], "en_flota": bool(v), "vid": v.get("id"),
                      "modelo": f"{v.get('marca') or ''} {v.get('modelo') or ''}".strip(),
                      "asientos": v.get("asientos") or 0})
    lista.sort(key=lambda c: (-c["servicios"], len(c["coche"]), c["coche"]))
    return lista


def _por_tripulante(servs):
    out = {}
    for s in servs:
        for t in s["trips"]:
            k = _clave_persona(t)
            p = out.setdefault(k, {"clave": k, "nombre": t, "servicios": 0, "refuerzos": 0, "km": 0,
                                   "dias": set(), "coches": set(), "ultimo": ""})
            p["servicios"] += 1
            p["refuerzos"] += 1 if s["refuerzo"] else 0
            p["km"] += s["km"] or 0
            p["dias"].add(s["fecha"])
            if s["coche"]:
                p["coches"].add(s["coche"])
            p["ultimo"] = max(p["ultimo"], f"{s['fecha']} {s['hora']}")
    lista = [{**p, "dias": len(p["dias"]), "coches": len(p["coches"]), "km": round(p["km"])} for p in out.values()]
    lista.sort(key=lambda p: (-p["servicios"], p["clave"]))
    return lista


def _flota_base(conn, hasta):
    """Los coches que hacen nacionales: los que aparecieron alguna vez en el
    tráfico. Así los de turismo o corporativos no cuentan como 'parados'."""
    base, ultimo = set(), {}
    for r in conn.execute("""SELECT s.coche, MAX(d.fecha || ' ' || COALESCE(s.hora, '')) AS ult
                             FROM trafico_salidas s JOIN trafico_dias d ON d.id = s.dia_id
                             WHERE d.fecha <= ? AND s.coche <> '' GROUP BY s.coche""", (hasta,)).fetchall():
        c = limpiar(r["coche"])
        if c.isdigit():
            base.add(c)
            ultimo[c] = max(ultimo.get(c, ""), r["ult"] or "")
    return base, ultimo


def _estado_parados(conn, numeros, flota, fecha):
    """Por qué puede estar parado: OT abierta o fuera de servicio."""
    ids = {flota[n]["id"]: n for n in numeros if n in flota}
    estado = {n: [] for n in numeros}
    if not ids:
        return estado
    marcas = ",".join("?" * len(ids))
    try:
        for r in conn.execute(f"""SELECT id, vehiculo_id FROM ordenes_trabajo
                                  WHERE vehiculo_id IN ({marcas}) AND COALESCE(estado,'') <> 'cerrada'""",
                              list(ids)).fetchall():
            n = ids[r["vehiculo_id"]]
            if not any(e["tipo"] == "taller" for e in estado[n]):
                estado[n].append({"tipo": "taller", "txt": f"OT #{r['id']} abierta"})
    except Exception:
        pass
    try:
        for r in conn.execute(f"""SELECT vehiculo_id, motivo FROM fuera_servicio
                                  WHERE vehiculo_id IN ({marcas}) AND fecha_desde <= ?
                                  AND (fecha_hasta IS NULL OR fecha_hasta = '' OR fecha_hasta >= ?)""",
                              list(ids) + [fecha, fecha]).fetchall():
            estado[ids[r["vehiculo_id"]]].append({"tipo": "fuera", "txt": "Fuera de servicio" + (f": {r['motivo']}" if r["motivo"] else "")})
    except Exception:
        pass
    return estado


def _planificacion(conn, desde, hasta, servs):
    """Cómo se arma el tráfico: anticipación de la primera publicación, versiones
    por día, cambios después de publicar y salidas que se mandaron A CONFIRMAR."""
    primera_salida = {}
    for s in servs:
        if s["hora"]:
            primera_salida[s["fecha"]] = min(primera_salida.get(s["fecha"], "99:99"), s["hora"])
    dias = {}
    for r in conn.execute("""SELECT d.fecha, d.version, d.modificado, v.version AS v, v.publicado_el,
                                    CASE WHEN v.version = 1 THEN v.datos ELSE NULL END AS datos, v.cambios
                             FROM trafico_dias d LEFT JOIN trafico_versiones v ON v.dia_id = d.id
                             WHERE d.fecha >= ? AND d.fecha <= ?
                             ORDER BY d.fecha, v.version""", (desde, hasta)).fetchall():
        d = dias.setdefault(r["fecha"], {"fecha": r["fecha"], "versiones": 0, "cambios": 0, "primera": None,
                                         "a_confirmar": 0, "publicadas": 0})
        if r["v"] is None:
            continue
        d["versiones"] += 1
        if r["v"] == 1:
            d["primera"] = r["publicado_el"]
            try:
                foto = json.loads(r["datos"] or "{}")
                filas = [x for b in foto.get("bloques", []) for lado in ("ida", "regreso") for x in b.get(lado, [])]
                d["publicadas"] = len(filas)
                d["a_confirmar"] = sum(1 for x in filas if not limpiar(x.get("coche")).isdigit()
                                       or not (x.get("trip1") or x.get("trip2")))
            except Exception:
                pass
        else:
            try:
                d["cambios"] += len(json.loads(r["cambios"] or "[]"))
            except Exception:
                pass
    con_salidas = {s["fecha"] for s in servs}
    lista = []
    for f in sorted(con_salidas):
        d = dias.get(f) or {"fecha": f, "versiones": 0, "cambios": 0, "primera": None, "a_confirmar": 0, "publicadas": 0}
        anticip = None
        if d["primera"] and primera_salida.get(f):
            try:
                pub = datetime.fromisoformat(str(d["primera"])[:19])
                sale = datetime.fromisoformat(f"{f}T{primera_salida[f]}:00")
                anticip = round((sale - pub).total_seconds() / 3600, 1)
            except ValueError:
                anticip = None
        lista.append({"fecha": f, "versiones": d["versiones"], "cambios": d["cambios"],
                      "anticipacion": anticip, "a_confirmar": d["a_confirmar"], "publicadas": d["publicadas"],
                      "publicado": bool(d["versiones"])})
    publicados = [d for d in lista if d["publicado"]]
    anticip = [d["anticipacion"] for d in publicados if d["anticipacion"] is not None]
    tot_pub = sum(d["publicadas"] for d in publicados)
    return {
        "dias": lista,
        "dias_publicados": len(publicados),
        "dias_sin_publicar": len(lista) - len(publicados),
        "anticipacion_mediana": _redondo(statistics.median(anticip)) if anticip else None,
        "versiones_promedio": round(sum(d["versiones"] for d in publicados) / len(publicados), 1) if publicados else None,
        "cambios": sum(d["cambios"] for d in publicados),
        "a_confirmar_pct": round(sum(d["a_confirmar"] for d in publicados) / tot_pub * 100, 1) if tot_pub else None,
    }


def _tramos_para_distancias(servs, distancias):
    """Los tramos que aparecen (para cargar su distancia una sola vez)."""
    out = {}
    for s in servs:
        k = s["clave_tramo"]
        if not k:
            continue
        t = out.setdefault(k, {"clave": k, "nombre": nombre_tramo(k), "servicios": 0, "km": distancias.get(k)})
        t["servicios"] += 1
    return sorted(out.values(), key=lambda t: (t["km"] is not None, -t["servicios"], t["nombre"]))


# ════════════════════════════════════════════════════════════════════════════
# EL PANEL
# ════════════════════════════════════════════════════════════════════════════

def panel(desde, hasta, corredor=None):
    hoy = _hoy()
    hasta = min(hasta, hoy)
    if desde > hasta:
        desde = hasta
    n_dias = (date.fromisoformat(hasta) - date.fromisoformat(desde)).days + 1
    ant_hasta = (date.fromisoformat(desde) - timedelta(days=1)).isoformat()
    ant_desde = (date.fromisoformat(desde) - timedelta(days=n_dias)).isoformat()

    conn = get_connection()
    try:
        flota = _flota(conn)
        distancias = _distancias(conn)
        servs_todos = leer_servicios(conn, desde, hasta, flota, distancias)
        ant_todos = leer_servicios(conn, ant_desde, ant_hasta, flota, distancias)
        corredores_todos = sorted({s["corredor"] for s in servs_todos})
        if corredor:
            servs = [s for s in servs_todos if s["corredor"] == corredor]
            ant = [s for s in ant_todos if s["corredor"] == corredor]
        else:
            servs, ant = servs_todos, ant_todos
        base, ultimo_global = _flota_base(conn, hasta)
        usados = {s["coche"] for s in servs_todos if s["coche"]}
        parados_n = sorted(base - usados, key=lambda c: (ultimo_global.get(c, ""), c))
        estado = _estado_parados(conn, parados_n, flota, hasta)
        plan = _planificacion(conn, desde, hasta, servs)
    finally:
        conn.close()

    fechas = _rango(desde, hasta)
    con_trafico = sorted({s["fecha"] for s in servs})
    con_trafico_ant = sorted({s["fecha"] for s in ant})
    tot = _totales(servs, len(con_trafico))
    tot_ant = _totales(ant, len(con_trafico_ant))
    serie = _serie_diaria(servs, fechas, len(base))
    usos = [d["uso"] for d in serie if d["uso"] is not None]
    tot["uso_flota"] = round(sum(usos) / len(usos), 1) if usos else None
    tot["flota_base"] = len(base)
    tot["coches_por_dia"] = round(sum(d["coches"] for d in serie if d["con_trafico"]) / len(con_trafico), 1) if con_trafico else 0
    serie_ant = _serie_diaria(ant, _rango(ant_desde, ant_hasta), len(base))
    usos_ant = [d["uso"] for d in serie_ant if d["uso"] is not None]
    tot_ant["uso_flota"] = round(sum(usos_ant) / len(usos_ant), 1) if usos_ant else None
    tot_ant["coches_por_dia"] = round(sum(d["coches"] for d in serie_ant if d["con_trafico"]) / len(con_trafico_ant), 1) if con_trafico_ant else 0

    coches = _por_coche(servs, flota)
    tripulantes = _por_tripulante(servs)
    parados = [{"coche": c, "ultimo": ultimo_global.get(c, ""), "estado": estado.get(c, []),
                "vid": (flota.get(c) or {}).get("id"),
                "modelo": f"{(flota.get(c) or {}).get('marca') or ''} {(flota.get(c) or {}).get('modelo') or ''}".strip()}
               for c in parados_n]
    tramos = _tramos_para_distancias(servs_todos, distancias)
    deltas = _deltas(tot, tot_ant)

    return {
        "desde": desde, "hasta": hasta, "dias_periodo": n_dias,
        "anterior": {"desde": ant_desde, "hasta": ant_hasta, **tot_ant},
        "corredor": corredor or "", "corredores_disponibles": corredores_todos,
        "totales": tot, "deltas": deltas,
        "serie": serie,
        "corredores": _por_corredor(servs),
        "horario": _mapa_horario(servs, con_trafico),
        "coches": coches,
        "parados": parados,
        "tripulantes": tripulantes,
        "planificacion": plan,
        "tramos": tramos,
        "hay_datos": bool(servs),
        "primer_dia": _primer_dia(),
        "puede_editar": _puede_editar(),
    }


def _primer_dia():
    conn = get_connection()
    try:
        r = conn.execute("""SELECT MIN(d.fecha) AS f FROM trafico_dias d
                            WHERE EXISTS (SELECT 1 FROM trafico_salidas s WHERE s.dia_id = d.id)""").fetchone()
        return r["f"] if r else None
    finally:
        conn.close()


def detalle(tipo, clave, desde, hasta):
    """Los servicios de un coche, un tripulante, un corredor o un día."""
    hasta = min(hasta, _hoy())
    conn = get_connection()
    try:
        servs = leer_servicios(conn, desde, hasta)
    finally:
        conn.close()
    if tipo == "coche":
        lista = [s for s in servs if s["coche"] == clave]
    elif tipo == "tripulante":
        k = _clave_persona(clave)
        lista = [s for s in servs if any(_clave_persona(t) == k for t in s["trips"])]
    elif tipo == "corredor":
        lista = [s for s in servs if s["corredor"] == clave]
    elif tipo == "dia":
        lista = [s for s in servs if s["fecha"] == clave]
    else:
        return None
    lista.sort(key=lambda s: s["hora"] or "99:99")
    lista.sort(key=lambda s: s["fecha"], reverse=True)
    resumen = _totales(lista, len({s["fecha"] for s in lista}))
    return {"tipo": tipo, "clave": clave, "desde": desde, "hasta": hasta, "resumen": resumen,
            "servicios": [{k: s[k] for k in ("fecha", "hora", "lado", "corredor", "destino", "coche", "trips",
                                             "refuerzo", "km", "origen", "llega")} for s in lista]}


# ════════════════════════════════════════════════════════════════════════════
# EXCEL
# ════════════════════════════════════════════════════════════════════════════

def excel(desde, hasta, corredor=None):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    hasta = min(hasta, _hoy())
    conn = get_connection()
    try:
        flota = _flota(conn)
        servs = leer_servicios(conn, desde, hasta, flota)
    finally:
        conn.close()
    if corredor:
        servs = [s for s in servs if s["corredor"] == corredor]

    wb = Workbook()
    navy = PatternFill("solid", fgColor="1D344E")
    cebra = PatternFill("solid", fgColor="F4F6F9")
    blanco = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    titulo = Font(name="Calibri", bold=True, color="1D344E", size=15)
    sub = Font(name="Calibri", color="5B6B7F", size=10)
    linea = Border(bottom=Side(style="thin", color="E1E6ED"))

    def hoja(ws, nombre, encabezados, filas, anchos):
        ws.title = nombre
        ws["A1"] = "LA SANTANIANA · Servicios" + ("" if nombre == "Servicios" else f" · {nombre}")
        ws["A1"].font = titulo
        rango = f"Del {desde[8:10]}/{desde[5:7]}/{desde[:4]} al {hasta[8:10]}/{hasta[5:7]}/{hasta[:4]}"
        ws["A2"] = rango + (f" · Corredor {corredor}" if corredor else "") + " · Sacado del tráfico de cada día"
        ws["A2"].font = sub
        fila0 = 4
        for i, (enc, ancho) in enumerate(zip(encabezados, anchos), start=1):
            c = ws.cell(fila0, i, enc)
            c.fill, c.font = navy, blanco
            c.alignment = Alignment(horizontal="center", vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = ancho
        ws.row_dimensions[fila0].height = 22
        for r, fila in enumerate(filas, start=fila0 + 1):
            for i, val in enumerate(fila, start=1):
                c = ws.cell(r, i, val)
                c.border = linea
                if (r - fila0) % 2 == 0:
                    c.fill = cebra
                if isinstance(val, (int, float)) and encabezados[i - 1] != "Coche":
                    c.number_format = "#,##0"
        ws.freeze_panes = ws.cell(fila0 + 1, 1)
        if filas:
            ws.auto_filter.ref = f"A{fila0}:{get_column_letter(len(encabezados))}{fila0 + len(filas)}"

    filas = []
    for s in servs:
        d = date.fromisoformat(s["fecha"])
        filas.append([d.strftime("%d/%m/%Y"), DIAS_SEMANA[d.weekday()].capitalize(), s["hora"], s["corredor"],
                      "Ida" if s["lado"] == "ida" else "Regreso", s["destino"], "Sí" if s["refuerzo"] else "",
                      int(s["coche"]) if s["coche"] else "A CONFIRMAR", s["trips"][0] if s["trips"] else "",
                      s["trips"][1] if len(s["trips"]) > 1 else "", s["km"] or None, s["asientos"] or None])
    hoja(wb.active, "Servicios", ["Fecha", "Día", "Hora", "Corredor", "Sentido", "Tramo", "Refuerzo", "Coche",
                                  "Tripulante 1", "Tripulante 2", "Km", "Asientos"],
         filas, [12, 11, 8, 26, 10, 28, 10, 12, 26, 26, 9, 10])
    hoja(wb.create_sheet(), "Por tripulante", ["Tripulante", "Servicios", "Días", "Refuerzos", "Coches distintos", "Km"],
         [[p["nombre"], p["servicios"], p["dias"], p["refuerzos"], p["coches"], p["km"] or None]
          for p in _por_tripulante(servs)], [30, 11, 8, 11, 16, 10])
    hoja(wb.create_sheet(), "Por coche", ["Coche", "Modelo", "Servicios", "Días", "Refuerzos", "Corredor principal", "Km"],
         [[int(c["coche"]), c["modelo"], c["servicios"], c["dias"], c["refuerzos"], c["corredor"], c["km"] or None]
          for c in _por_coche(servs, flota)], [10, 22, 11, 8, 11, 26, 10])
    hoja(wb.create_sheet(), "Por corredor", ["Corredor", "Servicios", "Ida", "Regreso", "Refuerzos", "Coches", "Asientos", "Km"],
         [[c["corredor"], c["servicios"], c["ida"], c["regreso"], c["refuerzos"], c["coches"], c["asientos"] or None,
           c["km"] or None] for c in _por_corredor(servs)], [28, 11, 8, 10, 11, 9, 11, 10])
    salida = io.BytesIO()
    wb.save(salida)
    return salida.getvalue()


# ════════════════════════════════════════════════════════════════════════════
# DISTANCIAS
# ════════════════════════════════════════════════════════════════════════════

def guardar_distancia(clave, km, usuario=""):
    clave = str(clave or "").strip()
    if not clave or "|" not in clave:
        return False, "Tramo inválido"
    if km in ("", None):
        km = None
    else:
        try:
            km = float(str(km).replace(".", "").replace(",", ".")) if isinstance(km, str) else float(km)
        except ValueError:
            return False, "Escribí los km con números"
        if km <= 0 or km > 3000:
            return False, "Esa distancia no parece real"
    conn = get_connection()
    try:
        r = conn.execute("SELECT id FROM servicios_distancias WHERE clave=?", (clave,)).fetchone()
        if r:
            conn.execute("UPDATE servicios_distancias SET km=?, actualizado=?, actualizado_por=? WHERE id=?",
                         (km, _ahora(), usuario, r["id"]))
        else:
            conn.execute("INSERT INTO servicios_distancias (clave, km, actualizado, actualizado_por) VALUES (?,?,?,?)",
                         (clave, km, _ahora(), usuario))
        conn.commit()
        return True, "Guardado"
    finally:
        conn.close()


# ════════════════════════════════════════════════════════════════════════════
# ENDPOINTS
# ════════════════════════════════════════════════════════════════════════════

def _puede_ver():
    return session.get("rol") in ("admin", "taller", "auditor")


def _puede_editar():
    return session.get("rol") in ("admin", "taller")


def _periodo():
    hoy = _hoy()
    try:
        hasta = date.fromisoformat(request.args.get("hasta") or hoy).isoformat()
        desde = date.fromisoformat(request.args.get("desde") or
                                   (date.fromisoformat(hasta) - timedelta(days=29)).isoformat()).isoformat()
    except ValueError:
        return None, None
    if desde > hasta:
        desde, hasta = hasta, desde
    if (date.fromisoformat(hasta) - date.fromisoformat(desde)).days > MAX_DIAS:
        desde = (date.fromisoformat(hasta) - timedelta(days=MAX_DIAS)).isoformat()
    return desde, hasta


@bp_servicios.route("/api/servicios/panel", methods=["GET"])
def api_panel():
    if not _puede_ver():
        return jsonify({"error": "Sin permiso"}), 403
    desde, hasta = _periodo()
    if not desde:
        return jsonify({"error": "Fechas inválidas"}), 400
    return jsonify(panel(desde, hasta, limpiar(request.args.get("corredor")) or None))


@bp_servicios.route("/api/servicios/detalle", methods=["GET"])
def api_detalle():
    if not _puede_ver():
        return jsonify({"error": "Sin permiso"}), 403
    desde, hasta = _periodo()
    r = detalle(request.args.get("tipo", ""), request.args.get("clave", ""), desde, hasta) if desde else None
    return jsonify(r) if r else (jsonify({"error": "No encontrado"}), 404)


@bp_servicios.route("/api/servicios/excel", methods=["GET"])
def api_excel():
    if not _puede_ver():
        return jsonify({"error": "Sin permiso"}), 403
    desde, hasta = _periodo()
    if not desde:
        return jsonify({"error": "Fechas inválidas"}), 400
    contenido = excel(desde, hasta, limpiar(request.args.get("corredor")) or None)
    nombre = f"Servicios-{desde}-al-{hasta}.xlsx"
    return send_file(io.BytesIO(contenido), as_attachment=True, download_name=nombre,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@bp_servicios.route("/api/servicios/distancia", methods=["PUT"])
def api_distancia():
    if not _puede_editar():
        return jsonify({"ok": False, "msg": "Sin permiso"}), 403
    d = request.json or {}
    ok, msg = guardar_distancia(d.get("clave"), d.get("km"), session.get("nombre") or session.get("usuario") or "")
    return jsonify({"ok": ok, "msg": msg}), (200 if ok else 400)
