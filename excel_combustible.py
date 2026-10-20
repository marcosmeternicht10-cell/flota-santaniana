"""Excel de las cargas de combustible de corporativos.

Un libro para sentarse a revisar, no un volcado de la tabla:

  Resumen      los números del período, por mes, dónde se carga y qué coches
               conviene mirar, con gráficos.
  Detalle      cada carga, con filtros; los totales de arriba siguen al filtro.
  Por coche    consumo real contra lo esperado para su motor.
  Por chofer   cuánto cargó cada uno y su consumo en los tramos propios.
  Por cliente  cuánto combustible se fue en cada empresa.

Los totales son fórmulas que apuntan al Detalle: si alguien corrige una carga
en el Excel, todo se recalcula. Además se guarda el resultado ya calculado en
cada fórmula, para que el archivo se vea completo aunque se abra en un visor
que no calcula (la vista previa de WhatsApp o del celular).
"""
import io
import re
import zipfile
import datetime

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.formatting.rule import CellIsRule, DataBarRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

NAVY = "1A2B4A"
AZUL = "1E5A96"
ROJO = "DC2641"
TINTA = "1D2939"
GRIS = "667085"
GRIS_CLARO = "98A2B3"
LINEA = "E4E7EC"
ZEBRA = "F7F9FC"
SUAVE = "F1F5FA"
FUENTE = "Arial"

F_GS = '#,##0;-#,##0;"–"'
F_ENTERO = '#,##0;-#,##0;"–"'
F_LITROS = '#,##0.0;-#,##0.0;"–"'
F_LITROS_R = '#,##0;-#,##0;"–"'     # en los resúmenes, sin decimales
F_L100 = '0.0;-0.0;"–"'
F_PCT = '0.0%;-0.0%;"–"'
F_DESVIO = '+0.0%;-0.0%;0.0%'
F_FECHA = "dd/mm/yyyy"
F_MES = "mmm yyyy"

MESES = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio",
         "agosto", "septiembre", "octubre", "noviembre", "diciembre"]


def _etiquetas():
    """Etiquetas de datos con solo el número (sin categoría ni serie)."""
    d = DataLabelList()
    d.showVal = True
    d.showCatName = False
    d.showSerName = False
    d.showLegendKey = False
    d.showPercent = False
    d.numFmt = "#,##0"
    return d


def _f(size=10, bold=False, color=TINTA, italic=False):
    return Font(name=FUENTE, size=size, bold=bold, color=color, italic=italic)


def _relleno(color):
    return PatternFill("solid", fgColor=color)


def _num(v):
    try:
        return float(v or 0)
    except (TypeError, ValueError):
        return 0.0


def _fecha(v):
    try:
        return datetime.date.fromisoformat(str(v)[:10])
    except (TypeError, ValueError):
        return None


def _fecha_txt(v):
    d = _fecha(v)
    return d.strftime("%d/%m/%Y") if d else ""


def _bus(v):
    try:
        return int(str(v).strip())
    except (TypeError, ValueError):
        return str(v or "").strip() or "—"


class _Libro:
    """Un Workbook que además anota el resultado de cada fórmula."""

    def __init__(self):
        self.wb = Workbook()
        self.cache = {}

    def formula(self, ws, celda, formula, valor):
        ws[celda] = formula
        self.cache[(ws.title, celda)] = valor
        return ws[celda]

    def guardar(self):
        buf = io.BytesIO()
        self.wb.calculation.fullCalcOnLoad = True
        self.wb.save(buf)
        return _poner_resultados(buf.getvalue(), self.wb, self.cache)


def _poner_resultados(contenido, wb, cache):
    """Escribe en el XML el resultado ya calculado de cada fórmula.

    openpyxl guarda las fórmulas sin resultado; Excel las calcula al abrir,
    pero los visores (celular, WhatsApp, vista previa del correo) muestran la
    celda vacía. Con el resultado puesto, se ve completo en cualquier lado.
    """
    por_hoja = {}
    for (hoja, celda), valor in cache.items():
        por_hoja.setdefault(hoja, {})[celda] = valor
    archivos = {f"xl/worksheets/sheet{i + 1}.xml": ws.title for i, ws in enumerate(wb.worksheets)}
    entrada = zipfile.ZipFile(io.BytesIO(contenido))
    salida_buf = io.BytesIO()
    with zipfile.ZipFile(salida_buf, "w", zipfile.ZIP_DEFLATED) as salida:
        for item in entrada.infolist():
            datos = entrada.read(item.filename)
            hoja = archivos.get(item.filename)
            if hoja and hoja in por_hoja:
                valores = por_hoja[hoja]
                xml = datos.decode("utf-8")

                def reemplazo(m):
                    celda, attrs, f = m.group(1), m.group(2), m.group(3)
                    if celda not in valores:
                        return m.group(0)
                    v = valores[celda]
                    if v is None or v == "":
                        return f'<c r="{celda}"{attrs} t="str">{f}<v></v></c>'
                    if isinstance(v, (int, float)) and not isinstance(v, bool):
                        return f'<c r="{celda}"{attrs}>{f}<v>{repr(float(v)) if isinstance(v, float) else v}</v></c>'
                    texto = (str(v).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
                    return f'<c r="{celda}"{attrs} t="str">{f}<v>{texto}</v></c>'

                xml = re.sub(r'<c r="([A-Z]+[0-9]+)"([^>]*)>(<f>.*?</f>)<v></v></c>', reemplazo, xml)
                datos = xml.encode("utf-8")
            salida.writestr(item, datos)
    return salida_buf.getvalue()


# ─── Datos ────────────────────────────────────────────────────────────────────

def _filas(cargas):
    """Cada carga como una fila del Detalle, ya con los textos que se ven."""
    filas = []
    for c in cargas:
        f = _fecha(c.get("fecha"))
        valido = bool(c.get("consumo_valido"))
        lugar = (c.get("lugar_carga") or "").strip()
        emblema = (c.get("emblema") or "").strip().upper()
        if lugar == "taller":
            lugar_txt, punto = "Taller", "Taller"
        elif lugar == "tercerizado":
            lugar_txt, punto = "Afuera", emblema or "Afuera (sin emblema)"
        else:
            lugar_txt, punto = "Sin dato", "Sin dato"
        if c.get("parcial_deducida"):
            estado = "Carga parcial (se suma a la siguiente)"
        elif c.get("rendimiento_dudoso"):
            estado = "Revisar kilometraje"
        elif c.get("sin_odometro"):
            estado = "Sin kilometraje"
        elif valido:
            estado = "Medida"
        elif not c.get("tanque_lleno", 1):
            estado = "Carga parcial"
        else:
            estado = "Primera carga"
        km = c.get("km_recorridos") if valido else None
        lt = c.get("litros_del_tramo") if valido else None
        filas.append({
            "fecha": f,
            "mes": f.replace(day=1) if f else None,
            "bus": _bus(c.get("n_interno") or c.get("patente")),
            "vid": c.get("vehiculo_id"),
            "patente": c.get("patente") or "",
            "modelo": " ".join(x for x in (c.get("marca") or "", c.get("modelo") or "") if x).strip(),
            "chofer": (c.get("chofer") or "").strip() or "Sin chofer",
            "cliente": (c.get("cliente") or "").strip() or "Sin cliente",
            "lugar": lugar_txt,
            "punto": punto,
            "odometro": _num(c.get("odometro")) or None,
            "litros": round(_num(c.get("litros")), 2),
            "lleno": "Sí" if c.get("tanque_lleno", 1) else "No",
            "km": km,
            "litros_tramo": lt,
            "l100": (lt / km * 100) if km and lt else None,
            "propio": "Sí" if valido and c.get("tramo_limpio", True) else "No",
            "precio": _num(c.get("precio_litro")) or None,
            "total": round(_num(c.get("costo_total"))) or None,
            "estado": estado,
            "obs": c.get("observaciones") or "",
            "cargo": c.get("registrado_por") or "",
        })
    filas.sort(key=lambda x: (x["fecha"] or datetime.date.min), reverse=True)
    return filas


def _agrupar(filas, clave):
    g = {}
    for f in filas:
        d = g.setdefault(f[clave], {"clave": f[clave], "cargas": 0, "litros": 0.0, "total": 0.0,
                                    "km": 0.0, "litros_tramo": 0.0, "km_p": 0.0, "lt_p": 0.0,
                                    "fila": f})
        d["cargas"] += 1
        d["litros"] += f["litros"]
        d["total"] += f["total"] or 0
        d["km"] += f["km"] or 0
        d["litros_tramo"] += f["litros_tramo"] or 0
        if f["propio"] == "Sí":
            d["km_p"] += f["km"] or 0
            d["lt_p"] += f["litros_tramo"] or 0
    return sorted(g.values(), key=lambda x: -x["litros"])


# ─── Piezas de formato ────────────────────────────────────────────────────────

def _encabezado(ws, titulo, subtitulo, ultima_col, logo=None, generado="", col_titulo=None):
    """Logo, título, bajada y la franja de la marca (rojo y azul)."""
    ws.sheet_view.showGridLines = False
    ws.row_dimensions[1].height = 40
    ws.row_dimensions[2].height = 18
    ws.row_dimensions[3].height = 5
    forzar_col, col_titulo = col_titulo, 2
    if logo:
        try:
            from openpyxl.drawing.image import Image as XLImage
            img = XLImage(logo)
            alto = 46
            img.width, img.height = round(img.width * alto / img.height), alto
            ws.add_image(img, "A1")
        except Exception:
            col_titulo = 1
    else:
        col_titulo = 1
    if forzar_col:
        col_titulo = forzar_col
    c = ws.cell(row=1, column=col_titulo, value=titulo)
    c.font = _f(17, True, NAVY)
    c.alignment = Alignment(vertical="bottom")
    c = ws.cell(row=2, column=col_titulo, value=subtitulo)
    c.font = _f(10, color=GRIS)
    c.alignment = Alignment(vertical="top")
    if generado:
        c = ws.cell(row=2, column=ultima_col, value=generado)
        c.font = _f(8.5, color=GRIS_CLARO, italic=True)
        c.alignment = Alignment(horizontal="right", vertical="top")
    corte = max(1, round(ultima_col * 0.28))
    for j in range(1, ultima_col + 1):
        ws.cell(row=3, column=j).fill = _relleno(ROJO if j <= corte else NAVY)


def _cabecera_tabla(ws, fila, columnas, col0=1):
    for j, (nombre, ancho, _) in enumerate(columnas, col0):
        c = ws.cell(row=fila, column=j, value=nombre)
        c.font = _f(9, True, "FFFFFF")
        c.fill = _relleno(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = Border(right=Side(style="thin", color="2E4166"))
        ws.column_dimensions[get_column_letter(j)].width = ancho
    ws.row_dimensions[fila].height = 32


def _celda_dato(c, formato=None, zebra=False, alinear=None, negrita=False):
    c.font = _f(9.5, negrita)
    c.border = Border(bottom=Side(style="hair", color=LINEA))
    if zebra:
        c.fill = _relleno(ZEBRA)
    if formato:
        c.number_format = formato
    if alinear == "texto":
        c.alignment = Alignment(horizontal="left", indent=1, vertical="center")
    elif alinear:
        c.alignment = Alignment(horizontal=alinear, vertical="center")
    else:
        c.alignment = Alignment(vertical="center")


def _fila_total(ws, fila, col_desde, col_hasta):
    for j in range(col_desde, col_hasta + 1):
        c = ws.cell(row=fila, column=j)
        c.font = _f(9.5, True, NAVY)
        c.fill = _relleno(SUAVE)
        c.border = Border(top=Side(style="medium", color=NAVY))
    ws.row_dimensions[fila].height = 22


def _titulo_seccion(ws, celda, texto, nota=None):
    ws[celda] = texto.upper()
    ws[celda].font = _f(9, True, AZUL)
    if nota:
        f, c = ws[celda].row, ws[celda].column
        n = ws.cell(row=f + 1, column=c, value=nota)
        n.font = _f(8.5, color=GRIS)


def _impresion(ws, filas_titulo=None, horizontal=True):
    ws.page_setup.orientation = "landscape" if horizontal else "portrait"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = 0.5
    ws.page_margins.bottom = 0.6
    ws.oddFooter.left.text = "La Santaniana · Combustible de corporativos"
    ws.oddFooter.left.size = 8
    ws.oddFooter.right.text = "Hoja &P de &N"
    ws.oddFooter.right.size = 8
    if filas_titulo:
        ws.print_title_rows = filas_titulo


# ─── El libro ─────────────────────────────────────────────────────────────────

def generar_excel_cargas(cargas, desde=None, hasta=None, chofer=None, usuario="",
                         logo=None, referencia=None):
    """Arma el Excel. `referencia(vehiculo_id)` devuelve el consumo esperado
    de ese coche (dict con litros_100km) o None."""
    L = _Libro()
    wb = L.wb
    filas = _filas(cargas)
    n = len(filas)

    fechas = [f["fecha"] for f in filas if f["fecha"]]
    d1 = _fecha_txt(desde) or (min(fechas).strftime("%d/%m/%Y") if fechas else "el inicio")
    d2 = _fecha_txt(hasta) or (max(fechas).strftime("%d/%m/%Y") if fechas else "hoy")
    periodo = f"Del {d1} al {d2}"
    filtro = f" · Chofer: {chofer}" if chofer else " · Todos los choferes"
    ahora = datetime.datetime.now()
    generado = f"Generado el {ahora.strftime('%d/%m/%Y %H:%M')}" + (f" por {usuario}" if usuario else "")

    # ── Detalle (va primero en el armado: el resto apunta acá) ──────────────
    wsR = wb.active
    wsR.title = "Resumen"
    wsD = wb.create_sheet("Detalle")
    wsC = wb.create_sheet("Por coche")
    wsH = wb.create_sheet("Por chofer")
    wsK = wb.create_sheet("Por cliente")
    for ws, color in ((wsR, ROJO), (wsD, NAVY), (wsC, AZUL), (wsH, AZUL), (wsK, AZUL)):
        ws.sheet_properties.tabColor = color

    cols_d = [("Fecha", 11.5, F_FECHA), ("Mes", 10, F_MES), ("Bus", 8, "0"), ("Patente", 10, None),
              ("Modelo", 22, None), ("Chofer", 24, None), ("Cliente", 22, None), ("Lugar", 9, None),
              ("Punto de carga", 15, None), ("Odómetro", 11, F_ENTERO), ("Litros", 10, F_LITROS),
              ("Tanque lleno", 8.5, None), ("Km del tramo", 10, F_ENTERO), ("Litros del tramo", 10, F_LITROS),
              ("L cada 100 km", 9.5, F_L100), ("Cuenta para el chofer", 10, None),
              ("Gs. por litro", 10.5, F_GS), ("Total Gs.", 13.5, F_GS), ("Estado", 17, None),
              ("Observaciones", 26, None), ("Cargó", 12, None)]
    U = len(cols_d)
    H0 = 5                      # fila de los títulos de columna
    P = H0 + 1                  # primera fila de datos
    Z = H0 + max(n, 1)          # última fila de datos
    _encabezado(wsD, "Detalle de cargas", periodo + filtro, U, None, generado)
    _cabecera_tabla(wsD, H0, cols_d)

    def rango(col):
        return f"Detalle!${col}${P}:${col}${Z}"

    claves = ["fecha", "mes", "bus", "patente", "modelo", "chofer", "cliente", "lugar", "punto",
              "odometro", "litros", "lleno", "km", "litros_tramo", None, "propio", "precio", "total",
              "estado", "obs", "cargo"]
    centrados = {"Mes", "Bus", "Lugar", "Tanque lleno", "Cuenta para el chofer", "Fecha"}
    for i, f in enumerate(filas):
        r = P + i
        z = i % 2 == 1
        for j, (clave, (nombre, _, formato)) in enumerate(zip(claves, cols_d), 1):
            c = wsD.cell(row=r, column=j)
            if clave is None:          # L cada 100 km: se calcula en la hoja
                L.formula(wsD, c.coordinate, f'=IF(AND(M{r}>0,N{r}>0),N{r}/M{r}*100,"")',
                          f["l100"] if f["l100"] else "")
            else:
                c.value = f[clave]
            _celda_dato(c, formato, z, "center" if nombre in centrados else
                        "texto" if nombre in ("Estado", "Observaciones", "Cargó") else None)
        wsD.row_dimensions[r].height = 17
    if not filas:
        wsD.cell(row=P, column=1, value="No hay cargas de corporativos en este período.").font = _f(10, color=GRIS, italic=True)

    # Totales que siguen al filtro, arriba de la tabla (quedan siempre a la vista)
    tot_litros = sum(f["litros"] for f in filas)
    tot_gs = sum(f["total"] or 0 for f in filas)
    tot_km = sum(f["km"] or 0 for f in filas)
    tot_lt = sum(f["litros_tramo"] or 0 for f in filas)
    wsD.row_dimensions[4].height = 22
    if filas:
        L.formula(wsD, "A4", f'="Mostrando "&SUBTOTAL(103,A{P}:A{Z})&" de {n} cargas"',
                  f"Mostrando {n} de {n} cargas")
    else:
        wsD["A4"] = "Sin cargas en el período"
    wsD["A4"].font = _f(9, True, AZUL)
    wsD["I4"] = "Totales de lo que se ve:"
    wsD["I4"].font = _f(9, True, GRIS)
    wsD["I4"].alignment = Alignment(horizontal="right", vertical="center")
    for col, formula, valor, fmt in (
            ("K", f"=SUBTOTAL(109,K{P}:K{Z})", tot_litros, F_LITROS),
            ("M", f"=SUBTOTAL(109,M{P}:M{Z})", tot_km, F_ENTERO),
            ("N", f"=SUBTOTAL(109,N{P}:N{Z})", tot_lt, F_LITROS),
            ("O", f'=IFERROR(SUBTOTAL(109,N{P}:N{Z})/SUBTOTAL(109,M{P}:M{Z})*100,"")',
             (tot_lt / tot_km * 100) if tot_km else "", F_L100),
            ("R", f"=SUBTOTAL(109,R{P}:R{Z})", tot_gs, F_GS)):
        c = L.formula(wsD, f"{col}4", formula, valor)
        c.font = _f(9.5, True, NAVY)
        c.number_format = fmt
        c.fill = _relleno(SUAVE)
        c.alignment = Alignment(horizontal="right", vertical="center")
        c.border = Border(bottom=Side(style="medium", color=NAVY))

    if filas:
        tabla = Table(displayName="Cargas", ref=f"A{H0}:{get_column_letter(U)}{Z}")
        tabla.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False,
                                              showColumnStripes=False)
        wsD.add_table(tabla)
        wsD.conditional_formatting.add(
            f"K{P}:K{Z}", DataBarRule(start_type="num", start_value=0, end_type="max",
                                      color="9DB8D9", showValue=True))
        wsD.conditional_formatting.add(
            f"S{P}:S{Z}", FormulaRule(formula=[f'$S{P}="Revisar kilometraje"'],
                                      font=Font(name=FUENTE, bold=True, color=ROJO)))
        wsD.conditional_formatting.add(
            f"S{P}:S{Z}", FormulaRule(formula=[f'$S{P}="Sin kilometraje"'],
                                      font=Font(name=FUENTE, color="B54708")))
    wsD.freeze_panes = f"D{P}"
    _impresion(wsD, f"{H0}:{H0}")

    # ── Por coche ────────────────────────────────────────────────────────────
    cols_c = [("Bus", 8, "0"), ("Patente", 10, None), ("Modelo", 24, None), ("Cargas", 8, F_ENTERO),
              ("Litros", 11, F_LITROS_R), ("Total Gs.", 14, F_GS), ("Km medidos", 11, F_ENTERO),
              ("Litros medidos", 11, F_LITROS_R), ("Consumo real\nL cada 100 km", 12, F_L100),
              ("Esperado del motor\nL cada 100 km", 13, F_L100), ("Diferencia", 10, F_DESVIO),
              ("Lectura", 20, None)]
    _encabezado(wsC, "Consumo por coche",
                periodo + filtro + " · real contra lo esperado para su motor", len(cols_c), None, generado)
    _cabecera_tabla(wsC, H0, cols_c)
    buses = _agrupar(filas, "bus")
    lecturas = {}
    for i, b in enumerate(buses):
        r = P + i
        z = i % 2 == 1
        f0 = b["fila"]
        esperado = None
        if referencia and f0.get("vid"):
            try:
                ref = referencia(f0["vid"])
                esperado = float(ref["litros_100km"]) if ref and ref.get("litros_100km") else None
            except Exception:
                esperado = None
        real = b["litros_tramo"] / b["km"] * 100 if b["km"] else None
        desvio = (real / esperado - 1) if (real and esperado) else None
        if esperado is None:
            lectura = "Sin consumo de referencia"
        elif real is None:
            lectura = "Todavía no se puede medir"
        elif desvio > 0.15:
            lectura = "Consume de más"
        elif desvio > 0.07:
            lectura = "Algo alto"
        elif desvio < -0.15:
            lectura = "Revisar kilometraje"
        else:
            lectura = "Dentro de lo esperado"
        lecturas[b["clave"]] = (lectura, desvio, real, esperado)
        bus = b["clave"]
        crit = f"A{r}"
        vals = [bus, f0["patente"], f0["modelo"]]
        for j, v in enumerate(vals, 1):
            wsC.cell(row=r, column=j, value=v)
        L.formula(wsC, f"D{r}", f"=COUNTIFS({rango('C')},{crit})", b["cargas"])
        L.formula(wsC, f"E{r}", f"=SUMIFS({rango('K')},{rango('C')},{crit})", round(b["litros"], 2))
        L.formula(wsC, f"F{r}", f"=SUMIFS({rango('R')},{rango('C')},{crit})", b["total"])
        L.formula(wsC, f"G{r}", f"=SUMIFS({rango('M')},{rango('C')},{crit})", b["km"])
        L.formula(wsC, f"H{r}", f"=SUMIFS({rango('N')},{rango('C')},{crit})", round(b["litros_tramo"], 2))
        L.formula(wsC, f"I{r}", f'=IF(G{r}>0,H{r}/G{r}*100,"")', real if real else "")
        wsC.cell(row=r, column=10, value=esperado)
        L.formula(wsC, f"K{r}", f'=IF(AND(ISNUMBER(I{r}),J{r}>0),I{r}/J{r}-1,"")',
                  desvio if desvio is not None else "")
        L.formula(wsC, f"L{r}",
                  f'=IF(NOT(J{r}>0),"Sin consumo de referencia",IF(NOT(ISNUMBER(I{r})),"Todavía no se puede medir",'
                  f'IF(K{r}>0.15,"Consume de más",IF(K{r}>0.07,"Algo alto",'
                  f'IF(K{r}<-0.15,"Revisar kilometraje","Dentro de lo esperado")))))', lectura)
        for j, (nombre, _, formato) in enumerate(cols_c, 1):
            _celda_dato(wsC.cell(row=r, column=j), formato, z,
                        "center" if nombre in ("Bus", "Cargas") else "texto" if nombre == "Lectura" else None,
                        negrita=(j == 1))
        wsC.cell(row=r, column=10).font = _f(9.5, color="1849A9")   # dato que viene del plan
        wsC.row_dimensions[r].height = 18
    zc = P + max(len(buses), 1) - 1
    tc = zc + 1
    if buses:
        _fila_total(wsC, tc, 1, len(cols_c))
        wsC.cell(row=tc, column=1, value="Total")
        L.formula(wsC, f"D{tc}", f"=SUM(D{P}:D{zc})", n)
        L.formula(wsC, f"E{tc}", f"=SUM(E{P}:E{zc})", round(tot_litros, 2))
        L.formula(wsC, f"F{tc}", f"=SUM(F{P}:F{zc})", tot_gs)
        L.formula(wsC, f"G{tc}", f"=SUM(G{P}:G{zc})", tot_km)
        L.formula(wsC, f"H{tc}", f"=SUM(H{P}:H{zc})", round(tot_lt, 2))
        L.formula(wsC, f"I{tc}", f'=IF(G{tc}>0,H{tc}/G{tc}*100,"")',
                  (tot_lt / tot_km * 100) if tot_km else "")
        for col, fmt in (("D", F_ENTERO), ("E", F_LITROS_R), ("F", F_GS), ("G", F_ENTERO),
                         ("H", F_LITROS_R), ("I", F_L100)):
            wsC[f"{col}{tc}"].number_format = fmt
        rng = f"K{P}:K{zc}"
        wsC.conditional_formatting.add(rng, CellIsRule(operator="greaterThan", formula=["0.15"],
                                                       fill=_relleno("FDE4E7"), font=Font(name=FUENTE, bold=True, color="B42318")))
        wsC.conditional_formatting.add(rng, CellIsRule(operator="between", formula=["0.0701", "0.15"],
                                                       fill=_relleno("FEF0C7"), font=Font(name=FUENTE, bold=True, color="B54708")))
        wsC.conditional_formatting.add(rng, CellIsRule(operator="lessThan", formula=["-0.15"],
                                                       fill=_relleno("EAECF0"), font=Font(name=FUENTE, bold=True, color="475467")))
        wsC.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(ISNUMBER(K{P}),K{P}>=-0.15,K{P}<=0.07)'],
                                                        fill=_relleno("DCFAE6"), font=Font(name=FUENTE, bold=True, color="067647")))
        wsC.conditional_formatting.add(f"E{P}:E{zc}", DataBarRule(start_type="num", start_value=0,
                                                                  end_type="max", color="9DB8D9", showValue=True))
        wsC.auto_filter.ref = f"A{H0}:{get_column_letter(len(cols_c))}{zc}"
        nota = tc + 2
        for k, t in enumerate((
                "Cómo leerlo",
                "· El consumo se mide de tanque lleno a tanque lleno: km entre las dos cargas y litros que hizo falta para volver a llenar.",
                "· Esperado del motor (en azul) sale del plan de mantenimiento del coche o del consumo propio que se le haya cargado.",
                "· Diferencia: más de 15% arriba = consume de más; entre 7% y 15% = algo alto; más de 15% abajo = seguramente falta cargar kilómetros.")):
            c = wsC.cell(row=nota + k, column=1, value=t)
            c.font = _f(8.5, k == 0, NAVY if k == 0 else GRIS)
    wsC.freeze_panes = f"B{P}"
    _impresion(wsC, f"{H0}:{H0}")

    # ── Por chofer ───────────────────────────────────────────────────────────
    cols_h = [("Chofer", 26, None), ("Cargas", 8, F_ENTERO), ("Litros", 11, F_LITROS_R),
              ("% de los litros", 10, F_PCT), ("Total Gs.", 14, F_GS), ("Km medidos\n(tramos propios)", 12, F_ENTERO),
              ("Litros medidos\n(tramos propios)", 12, F_LITROS_R), ("Consumo\nL cada 100 km", 11, F_L100),
              ("Coches que usó", 10, F_ENTERO)]
    _encabezado(wsH, "Cargas por chofer", periodo + filtro, len(cols_h), None, generado)
    _cabecera_tabla(wsH, H0, cols_h)
    choferes = _agrupar(filas, "chofer")
    for i, h in enumerate(choferes):
        r = P + i
        z = i % 2 == 1
        crit = f"A{r}"
        wsH.cell(row=r, column=1, value=h["clave"])
        L.formula(wsH, f"B{r}", f"=COUNTIFS({rango('F')},{crit})", h["cargas"])
        L.formula(wsH, f"C{r}", f"=SUMIFS({rango('K')},{rango('F')},{crit})", round(h["litros"], 2))
        L.formula(wsH, f"D{r}", f"=IFERROR(C{r}/SUM({rango('K')}),0)",
                  h["litros"] / tot_litros if tot_litros else 0)
        L.formula(wsH, f"E{r}", f"=SUMIFS({rango('R')},{rango('F')},{crit})", h["total"])
        L.formula(wsH, f"F{r}", f'=SUMIFS({rango("M")},{rango("F")},{crit},{rango("P")},"Sí")', h["km_p"])
        L.formula(wsH, f"G{r}", f'=SUMIFS({rango("N")},{rango("F")},{crit},{rango("P")},"Sí")', round(h["lt_p"], 2))
        L.formula(wsH, f"H{r}", f'=IF(F{r}>0,G{r}/F{r}*100,"")', (h["lt_p"] / h["km_p"] * 100) if h["km_p"] else "")
        wsH.cell(row=r, column=9, value=len({f["bus"] for f in filas if f["chofer"] == h["clave"]}))
        for j, (nombre, _, formato) in enumerate(cols_h, 1):
            _celda_dato(wsH.cell(row=r, column=j), formato, z,
                        "center" if nombre in ("Cargas", "Coches que usó") else None, negrita=(j == 1))
        wsH.row_dimensions[r].height = 18
    zh = P + max(len(choferes), 1) - 1
    th = zh + 1
    if choferes:
        _fila_total(wsH, th, 1, len(cols_h))
        wsH.cell(row=th, column=1, value="Total")
        km_p = sum(h["km_p"] for h in choferes)
        lt_p = sum(h["lt_p"] for h in choferes)
        for col, formula, valor, fmt in (
                ("B", f"=SUM(B{P}:B{zh})", n, F_ENTERO),
                ("C", f"=SUM(C{P}:C{zh})", round(tot_litros, 2), F_LITROS_R),
                ("D", f"=SUM(D{P}:D{zh})", 1 if tot_litros else 0, F_PCT),
                ("E", f"=SUM(E{P}:E{zh})", tot_gs, F_GS),
                ("F", f"=SUM(F{P}:F{zh})", km_p, F_ENTERO),
                ("G", f"=SUM(G{P}:G{zh})", round(lt_p, 2), F_LITROS_R),
                ("H", f'=IF(F{th}>0,G{th}/F{th}*100,"")', (lt_p / km_p * 100) if km_p else "", F_L100)):
            L.formula(wsH, f"{col}{th}", formula, valor).number_format = fmt
        wsH.conditional_formatting.add(f"C{P}:C{zh}", DataBarRule(start_type="num", start_value=0,
                                                                  end_type="max", color="9DB8D9", showValue=True))
        wsH.auto_filter.ref = f"A{H0}:{get_column_letter(len(cols_h))}{zh}"
        c = wsH.cell(row=th + 2, column=1,
                     value="El consumo de cada chofer cuenta solo los tramos que cargó él de punta a punta: "
                           "si en el medio hubo una carga sin kilometraje o de otro chofer, ese tramo no se le atribuye.")
        c.font = _f(8.5, color=GRIS)
    wsH.freeze_panes = f"B{P}"
    _impresion(wsH, f"{H0}:{H0}", horizontal=False)

    # ── Por cliente ──────────────────────────────────────────────────────────
    cols_k = [("Cliente", 28, None), ("Cargas", 8, F_ENTERO), ("Litros", 11, F_LITROS_R),
              ("% de los litros", 10, F_PCT), ("Total Gs.", 14, F_GS), ("Coches", 9, F_ENTERO),
              ("Choferes", 9, F_ENTERO)]
    _encabezado(wsK, "Combustible por cliente", periodo + filtro, len(cols_k), None, generado)
    _cabecera_tabla(wsK, H0, cols_k)
    clientes = _agrupar(filas, "cliente")
    for i, k in enumerate(clientes):
        r = P + i
        z = i % 2 == 1
        crit = f"A{r}"
        wsK.cell(row=r, column=1, value=k["clave"])
        L.formula(wsK, f"B{r}", f"=COUNTIFS({rango('G')},{crit})", k["cargas"])
        L.formula(wsK, f"C{r}", f"=SUMIFS({rango('K')},{rango('G')},{crit})", round(k["litros"], 2))
        L.formula(wsK, f"D{r}", f"=IFERROR(C{r}/SUM({rango('K')}),0)", k["litros"] / tot_litros if tot_litros else 0)
        L.formula(wsK, f"E{r}", f"=SUMIFS({rango('R')},{rango('G')},{crit})", k["total"])
        del_cliente = [f for f in filas if f["cliente"] == k["clave"]]
        wsK.cell(row=r, column=6, value=len({f["bus"] for f in del_cliente}))
        wsK.cell(row=r, column=7, value=len({f["chofer"] for f in del_cliente}))
        for j, (nombre, _, formato) in enumerate(cols_k, 1):
            _celda_dato(wsK.cell(row=r, column=j), formato, z,
                        "center" if nombre in ("Cargas", "Coches", "Choferes") else None, negrita=(j == 1))
        wsK.row_dimensions[r].height = 18
    zk = P + max(len(clientes), 1) - 1
    tk = zk + 1
    if clientes:
        _fila_total(wsK, tk, 1, len(cols_k))
        wsK.cell(row=tk, column=1, value="Total")
        for col, formula, valor, fmt in (
                ("B", f"=SUM(B{P}:B{zk})", n, F_ENTERO),
                ("C", f"=SUM(C{P}:C{zk})", round(tot_litros, 2), F_LITROS_R),
                ("D", f"=SUM(D{P}:D{zk})", 1 if tot_litros else 0, F_PCT),
                ("E", f"=SUM(E{P}:E{zk})", tot_gs, F_GS)):
            L.formula(wsK, f"{col}{tk}", formula, valor).number_format = fmt
        wsK.conditional_formatting.add(f"C{P}:C{zk}", DataBarRule(start_type="num", start_value=0,
                                                                  end_type="max", color="9DB8D9", showValue=True))
    wsK.freeze_panes = f"B{P}"
    _impresion(wsK, f"{H0}:{H0}", horizontal=False)

    # ── Resumen ──────────────────────────────────────────────────────────────
    ws = wsR
    ws.column_dimensions["A"].width = 2
    for j in range(2, 14):
        ws.column_dimensions[get_column_letter(j)].width = 12.5
    ws.column_dimensions["N"].width = 2
    _encabezado(ws, "Combustible de corporativos", periodo + filtro, 13, None, generado, col_titulo=2)
    # Sin logo en la columna A angosta: va a la derecha del título
    if logo:
        try:
            from openpyxl.drawing.image import Image as XLImage
            img = XLImage(logo)
            alto = 46
            img.width, img.height = round(img.width * alto / img.height), alto
            ws.add_image(img, "M1")
        except Exception:
            pass
    ws["B1"].alignment = Alignment(vertical="bottom")

    tramos = sum(1 for f in filas if f["km"])
    con_precio = [f for f in filas if f["precio"]]
    lit_precio = sum(f["litros"] for f in con_precio)
    gs_precio = sum(f["total"] or 0 for f in con_precio)
    n_buses = len({f["bus"] for f in filas})
    n_choferes = len({f["chofer"] for f in filas})
    dias = len({f["fecha"] for f in filas if f["fecha"]})
    afuera = sum(f["litros"] for f in filas if f["lugar"] == "Afuera")

    _titulo_seccion(ws, "B5", "Los números del período")
    tarjetas = [
        ("Cargas", f"=COUNT({rango('A')})", n, F_ENTERO, f"en {dias} día{'s' if dias != 1 else ''} con cargas"),
        ("Litros cargados", f"=SUM({rango('K')})", round(tot_litros, 2), F_LITROS_R, "todo lo que se cargó"),
        ("Total en guaraníes", f"=SUM({rango('R')})", tot_gs, F_GS, "litros × precio de cada carga"),
        ("Precio promedio por litro", f'=IFERROR(SUMIFS({rango("R")},{rango("Q")},">0")/SUMIFS({rango("K")},{rango("Q")},">0"),0)',
         (gs_precio / lit_precio) if lit_precio else 0, F_GS, "pesado por los litros de cada carga"),
        ("Consumo promedio", f'=IFERROR(SUM({rango("N")})/SUM({rango("M")})*100,0)',
         (tot_lt / tot_km * 100) if tot_km else 0, '0.0" L/100 km";-0.0;"–"',
         f"sobre {tramos} tramo{'s' if tramos != 1 else ''} medido{'s' if tramos != 1 else ''} tanque a tanque"),
        ("Km medidos", f"=SUM({rango('M')})", tot_km, F_ENTERO, "entre cargas con kilometraje"),
        ("Coches", f"=SUMPRODUCT(({rango('C')}<>\"\")/COUNTIF({rango('C')},{rango('C')}&\"\"))" if filas else "=0",
         n_buses, F_ENTERO, "distintos que cargaron"),
        ("Cargado afuera", f'=IFERROR(SUMIFS({rango("K")},{rango("H")},"Afuera")/SUM({rango("K")}),0)',
         (afuera / tot_litros) if tot_litros else 0, F_PCT, "de los litros, fuera del taller"),
    ]
    for t, (nombre, formula, valor, fmt, nota) in enumerate(tarjetas):
        fila = 7 + (t // 4) * 5
        col = 2 + (t % 4) * 3
        letra = get_column_letter(col)
        fin = get_column_letter(col + 2)
        for rr in range(fila, fila + 4):
            for cc in range(col, col + 3):
                celda = ws.cell(row=rr, column=cc)
                celda.fill = _relleno(SUAVE)
                izq = Side(style="thick", color=ROJO if t == 0 else AZUL) if cc == col else None
                celda.border = Border(left=izq) if izq else Border()
        ws.merge_cells(f"{letra}{fila + 1}:{fin}{fila + 2}")
        c = ws.cell(row=fila, column=col, value=nombre)
        c.font = _f(9, True, GRIS)
        c.alignment = Alignment(horizontal="left", indent=1, vertical="bottom")
        c = L.formula(ws, f"{letra}{fila + 1}", formula, valor)
        c.font = _f(20, True, NAVY)
        c.number_format = fmt
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        c = ws.cell(row=fila + 3, column=col, value=nota)
        c.font = _f(8.5, color=GRIS)
        c.alignment = Alignment(horizontal="left", indent=1, vertical="top")
        for rr, alto in ((fila, 20), (fila + 1, 16), (fila + 2, 16), (fila + 3, 18)):
            ws.row_dimensions[rr].height = alto
    ws.row_dimensions[11].height = 10

    # Por mes (izquierda) y dónde se carga (derecha)
    fila_t = 18
    _titulo_seccion(ws, f"B{fila_t}", "Mes a mes")
    _titulo_seccion(ws, f"I{fila_t}", "Dónde se carga")
    fila_c = fila_t + 1
    cab_mes = [("Mes", None, None), ("Cargas", None, None), ("Litros", None, None), ("Total Gs.", None, None)]
    for j, (nombre, _, _) in enumerate(cab_mes):
        col = 2 + (0 if j == 0 else j + 1)   # Mes ocupa B:C
        c = ws.cell(row=fila_c, column=col, value=nombre)
    ws.merge_cells(start_row=fila_c, start_column=2, end_row=fila_c, end_column=3)
    for col in range(2, 7):
        c = ws.cell(row=fila_c, column=col)
        c.font = _f(9, True, "FFFFFF")
        c.fill = _relleno(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
    meses = {}
    for f in filas:
        if f["mes"]:
            m = meses.setdefault(f["mes"], [0, 0.0, 0.0])
            m[0] += 1
            m[1] += f["litros"]
            m[2] += f["total"] or 0
    lista_meses = sorted(meses.items())
    for i, (mes, (cg, lt, gs)) in enumerate(lista_meses):
        r = fila_c + 1 + i
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=3)
        ws.cell(row=r, column=2, value=mes)
        L.formula(ws, f"D{r}", f"=COUNTIFS({rango('B')},B{r})", cg)
        L.formula(ws, f"E{r}", f"=SUMIFS({rango('K')},{rango('B')},B{r})", round(lt, 2))
        L.formula(ws, f"F{r}", f"=SUMIFS({rango('R')},{rango('B')},B{r})", gs)
        for col, fmt in ((2, F_MES), (3, None), (4, F_ENTERO), (5, F_LITROS_R), (6, F_GS)):
            _celda_dato(ws.cell(row=r, column=col), fmt, i % 2 == 1, "center" if col in (2, 4) else None)
        ws.cell(row=r, column=2).font = _f(9.5, True)
    fin_meses = fila_c + len(lista_meses)

    cab_pto = ["Punto de carga", "", "Cargas", "Litros", "%"]
    ws.merge_cells(start_row=fila_c, start_column=9, end_row=fila_c, end_column=10)
    for j, nombre in enumerate(cab_pto):
        c = ws.cell(row=fila_c, column=9 + j, value=nombre or None)
        c.font = _f(9, True, "FFFFFF")
        c.fill = _relleno(NAVY)
        c.alignment = Alignment(horizontal="center", vertical="center")
    puntos = _agrupar(filas, "punto")
    for i, p in enumerate(puntos):
        r = fila_c + 1 + i
        ws.merge_cells(start_row=r, start_column=9, end_row=r, end_column=10)
        ws.cell(row=r, column=9, value=p["clave"])
        L.formula(ws, f"K{r}", f"=COUNTIFS({rango('I')},I{r})", p["cargas"])
        L.formula(ws, f"L{r}", f"=SUMIFS({rango('K')},{rango('I')},I{r})", round(p["litros"], 2))
        L.formula(ws, f"M{r}", f"=IFERROR(L{r}/SUM({rango('K')}),0)", p["litros"] / tot_litros if tot_litros else 0)
        for col, fmt in ((9, None), (10, None), (11, F_ENTERO), (12, F_LITROS_R), (13, F_PCT)):
            _celda_dato(ws.cell(row=r, column=col), fmt, i % 2 == 1, "center" if col == 11 else None)
        ws.cell(row=r, column=9).font = _f(9.5, True)
    fin_puntos = fila_c + len(puntos)
    ws.row_dimensions[fila_c].height = 22

    # Gráficos debajo de cada tabla
    fila_g = max(fin_meses, fin_puntos) + 3
    for celda, texto in ((f"B{fila_g - 1}", "Litros por mes" if lista_meses else None),
                         (f"I{fila_g - 1}", "Litros por punto de carga" if puntos else None)):
        if texto:
            ws[celda] = texto
            ws[celda].font = _f(9, True, GRIS)
    if lista_meses:
        g = BarChart()
        g.type = "col"
        g.title = None
        g.style = 10
        g.add_data(Reference(ws, min_col=5, min_row=fila_c, max_row=fin_meses), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=2, min_row=fila_c + 1, max_row=fin_meses))
        g.legend = None
        g.gapWidth = 70
        g.y_axis.numFmt = "#,##0"
        g.y_axis.majorGridlines.spPr = None
        g.x_axis.number_format = F_MES
        g.x_axis.delete = False
        g.y_axis.delete = False
        s = g.series[0]
        s.graphicalProperties.solidFill = AZUL
        s.graphicalProperties.line.solidFill = AZUL
        s.dLbls = _etiquetas()
        g.height, g.width = 6.2, 11.6
        ws.add_chart(g, f"B{fila_g}")
    if puntos:
        g = BarChart()
        g.type = "bar"
        g.title = None
        g.style = 10
        g.add_data(Reference(ws, min_col=12, min_row=fila_c, max_row=fin_puntos), titles_from_data=True)
        g.set_categories(Reference(ws, min_col=9, min_row=fila_c + 1, max_row=fin_puntos))
        g.legend = None
        g.gapWidth = 60
        g.x_axis.scaling.orientation = "maxMin"   # el más grande arriba
        g.y_axis.numFmt = "#,##0"
        g.y_axis.majorGridlines = None
        g.x_axis.delete = False
        g.y_axis.delete = True
        s = g.series[0]
        s.graphicalProperties.solidFill = NAVY
        s.graphicalProperties.line.solidFill = NAVY
        s.dLbls = _etiquetas()
        g.height, g.width = 6.2, 11.6
        ws.add_chart(g, f"I{fila_g}")

    # Coches para mirar: los que se salen de lo esperado
    fila_m = fila_g + 14
    _titulo_seccion(ws, f"B{fila_m}", "Coches para mirar",
                    "Los que consumen fuera de lo esperado para su motor. El detalle completo está en la hoja Por coche.")
    cab = [("Bus", 2, 2), ("Modelo", 3, 5), ("Real L/100 km", 6, 7), ("Esperado", 8, 9),
           ("Diferencia", 10, 10), ("Lectura", 11, 13)]
    fc = fila_m + 3
    for nombre, c1, c2 in cab:
        if c2 > c1:
            ws.merge_cells(start_row=fc, start_column=c1, end_row=fc, end_column=c2)
        c = ws.cell(row=fc, column=c1, value=nombre)
        for cc in range(c1, c2 + 1):
            ws.cell(row=fc, column=cc).fill = _relleno(NAVY)
        c.font = _f(9, True, "FFFFFF")
        c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[fc].height = 22
    fila_bus = {b["clave"]: P + i for i, b in enumerate(buses)}
    mirar = [(bus, *lecturas[bus]) for bus in lecturas
             if lecturas[bus][0] in ("Consume de más", "Algo alto", "Revisar kilometraje")]
    mirar.sort(key=lambda x: -abs(x[2] or 0))
    if not mirar:
        c = ws.cell(row=fc + 1, column=2,
                    value="Ningún coche se sale de lo esperado en este período (o falta el consumo de referencia de sus motores).")
        c.font = _f(9.5, color=GRIS, italic=True)
    for i, (bus, lectura, desvio, real, esperado) in enumerate(mirar[:10]):
        r = fc + 1 + i
        rc = fila_bus[bus]
        for c1, c2 in ((3, 5), (6, 7), (8, 9), (11, 13)):
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
        f0 = next(b["fila"] for b in buses if b["clave"] == bus)
        L.formula(ws, f"B{r}", f"='Por coche'!A{rc}", bus)
        L.formula(ws, f"C{r}", f"='Por coche'!C{rc}", f0["modelo"])
        L.formula(ws, f"F{r}", f"='Por coche'!I{rc}", real)
        L.formula(ws, f"H{r}", f"='Por coche'!J{rc}", esperado)
        L.formula(ws, f"J{r}", f"='Por coche'!K{rc}", desvio)
        L.formula(ws, f"K{r}", f"='Por coche'!L{rc}", lectura)
        for col, fmt in ((2, "0"), (3, None), (6, F_L100), (8, F_L100), (10, F_DESVIO), (11, None)):
            _celda_dato(ws.cell(row=r, column=col), fmt, i % 2 == 1, "center" if col in (2, 6, 8, 10) else None)
        color = ("B42318", "FDE4E7") if lectura == "Consume de más" else \
                ("B54708", "FEF0C7") if lectura == "Algo alto" else ("475467", "EAECF0")
        for col in (10, 11):
            ws.cell(row=r, column=col).font = _f(9.5, True, color[0])
            ws.cell(row=r, column=col).fill = _relleno(color[1])
        ws.cell(row=r, column=2).font = _f(9.5, True)
    fin = fc + 1 + max(len(mirar[:10]), 1) + 1
    for k, t in enumerate((
            "Notas",
            "· Son las cargas de los choferes de corporativos (origen corporativo), filtradas por la fecha de carga.",
            "· El consumo se mide de tanque lleno a tanque lleno; las cargas parciales se suman al tramo y las que no tienen kilometraje no se pueden medir.",
            "· Los totales son fórmulas sobre la hoja Detalle: si se corrige una carga ahí, todo el libro se recalcula.")):
        c = ws.cell(row=fin + k, column=2, value=t)
        c.font = _f(8.5, k == 0, NAVY if k == 0 else GRIS)
    _impresion(ws, None, horizontal=False)
    ws.print_area = f"A1:N{fin + 4}"

    wb.active = 0
    return L.guardar()
